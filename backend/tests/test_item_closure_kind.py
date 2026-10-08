"""업무 종결 구분(item_closures): 저장·검증·응답·이벤트·엑셀 라벨·권한 불변·마이그레이션."""
import io
import re
import sqlite3
from pathlib import Path

import pytest
from alembic import command
from openpyxl import load_workbook
from sqlalchemy import select

from app.models.item_closure import ItemClosure
from tests.test_close_delete import REASON, call, events, get_item, meeting, team  # noqa: F401  (meeting·team 은 픽스처)
from tests.test_upload_processing import _alembic, env  # noqa: F401  (env 는 픽스처)

_BEFORE = "e5b1c7d29a64"  # item_closures 마이그레이션의 down_revision
_ROOT = Path(__file__).resolve().parent.parent


def close(env, who, item_id, **extra):
    return call(env, "POST", who, f"/api/action-items/{item_id}/close", json={**REASON, **extra})


def rows(env) -> list[ItemClosure]:
    with env["factory"]() as s:
        return list(s.scalars(select(ItemClosure).order_by(ItemClosure.id)))


@pytest.mark.parametrize("kind", ["completed", "forced"])
def test_kind_is_saved_with_response_and_event(env, team, meeting, kind):
    item_id = meeting[1][2]
    res = close(env, team["lead"], item_id, closureKind=kind)
    assert res.status_code == 200 and res.json()["status"] == "closed" and res.json()["closureKind"] == kind
    (row,) = rows(env)
    assert (row.item_id, row.kind, row.closed_by, row.tenant_id) == (item_id, kind, team["lead"].id, team["lead"].tenant_id)
    assert events(env, item_id)[-1].payload == {
        "before": {"status": "confirmed"}, "after": {"status": "closed"}, "reason": REASON["reason"], "closureKind": kind}
    # 회의록 상세에도 나온다
    detail = call(env, "GET", team["lead"], f"/api/meetings/{meeting[0]}").json()
    assert {i["id"]: i["closureKind"] for i in detail["actionItems"]}[item_id] == kind


@pytest.mark.parametrize("bad", ["done", "", "COMPLETED", 1])
def test_invalid_kind_is_422_and_nothing_changes(env, team, meeting, bad):
    item_id = meeting[1][2]
    assert close(env, team["lead"], item_id, closureKind=bad).status_code == 422
    assert get_item(env, item_id).status == "confirmed" and rows(env) == [] and events(env, item_id) == []


def test_without_kind_no_row_and_null_kind(env, team, meeting):
    item_id = meeting[1][2]
    res = close(env, team["lead"], item_id)
    assert res.status_code == 200 and res.json()["closureKind"] is None
    assert rows(env) == []
    assert "closureKind" not in events(env, item_id)[-1].payload
    # null 을 명시해도 같다
    other = meeting[1][1]
    assert close(env, team["lead"], other, closureKind=None).json()["closureKind"] is None and rows(env) == []


def test_reason_is_still_required_with_kind(env, team, meeting):
    res = call(env, "POST", team["lead"], f"/api/action-items/{meeting[1][2]}/close", json={"closureKind": "completed"})
    assert res.status_code == 422 and rows(env) == []


def test_repeat_call_keeps_first_kind_and_events(env, team, meeting):
    item_id = meeting[1][2]
    close(env, team["lead"], item_id, closureKind="completed")
    again = close(env, team["exe"], item_id, closureKind="forced")
    assert again.status_code == 200 and again.json()["closureKind"] == "completed"
    assert [r.kind for r in rows(env)] == ["completed"]
    assert [e.event_type for e in events(env, item_id)] == ["item.closed"]
    # 구분 없이 종결된 업무에 나중에 구분을 보내도 행을 만들지 않는다
    other = meeting[1][1]
    close(env, team["lead"], other)
    assert close(env, team["lead"], other, closureKind="completed").json()["closureKind"] is None
    assert [r.item_id for r in rows(env)] == [item_id]


def test_permissions_and_conditions_unchanged(env, team, meeting):
    _, (lead_item, mgr2_item, staff_item, pending_item) = meeting
    assert close(env, team["staff"], staff_item, closureKind="completed").status_code == 403
    assert close(env, team["mgr2"], staff_item, closureKind="completed").status_code == 403
    assert close(env, team["lead"], pending_item, closureKind="completed").status_code == 409
    assert close(env, team["mgr2"], mgr2_item, closureKind="forced").status_code == 200  # 본인 담당 업무
    assert [r.item_id for r in rows(env)] == [mgr2_item]


def test_export_status_labels(env, team, meeting):
    meeting_id, (lead_item, mgr2_item, staff_item, pending_item) = meeting
    close(env, team["lead"], lead_item, closureKind="completed")
    close(env, team["lead"], mgr2_item, closureKind="forced")
    close(env, team["lead"], staff_item)
    res = call(env, "GET", team["lead"], f"/api/meetings/{meeting_id}/export")
    sheet = load_workbook(io.BytesIO(res.content))["업무"]
    status = {r[0].value: r[5].value for r in sheet.iter_rows(min_row=2)}
    assert status == {lead_item: "완료", mgr2_item: "직권 종료", staff_item: "종결", pending_item: "확정 대기"}


def test_migration_roundtrip_and_downgrade_guard(tmp_path):
    url = f"sqlite:///{(tmp_path / 'm.db').as_posix()}"
    cfg = _alembic(url)
    command.upgrade(cfg, "head")
    command.downgrade(cfg, _BEFORE)
    command.upgrade(cfg, "head")
    command.check(cfg)
    raw = sqlite3.connect(tmp_path / "m.db")
    raw.execute("INSERT INTO item_closures (tenant_id, item_id, kind, closed_by, created_at) VALUES (1, 1, 'completed', 1, '2026-10-01')")
    raw.commit()
    raw.close()
    with pytest.raises(RuntimeError, match="downgrade 거부"):
        command.downgrade(cfg, _BEFORE)
    check = sqlite3.connect(tmp_path / "m.db")
    try:
        assert check.execute("SELECT COUNT(*) FROM item_closures").fetchone()[0] == 1
        with pytest.raises(sqlite3.IntegrityError):  # CHECK 제약
            check.execute("INSERT INTO item_closures (tenant_id, item_id, kind, closed_by, created_at) VALUES (1, 2, 'x', 1, '2026-10-01')")
        with pytest.raises(sqlite3.IntegrityError):  # 업무당 1행
            check.execute("INSERT INTO item_closures (tenant_id, item_id, kind, closed_by, created_at) VALUES (1, 1, 'forced', 1, '2026-10-01')")
    finally:
        check.close()


def test_alembic_offline_postgresql_ddl_has_closure_table():
    cfg = _alembic("postgresql+psycopg://u:p@localhost/db")
    buffer = io.StringIO()
    cfg.output_buffer = buffer
    command.upgrade(cfg, "head", sql=True)
    sql = buffer.getvalue()
    assert "CREATE TABLE item_closures" in sql
    assert "ck_item_closures_kind_valid" in sql and "uq_item_closures_item_id" in sql


def test_closure_table_is_not_used_outside_its_owners():
    allowed = {"models/item_closure.py", "models/__init__.py", "services/item_closure.py"}
    offenders = []
    for path in (_ROOT / "app").rglob("*.py"):
        rel = path.relative_to(_ROOT / "app").as_posix()
        if rel not in allowed and re.search(r"\bItemClosure\b|item_closures", path.read_text(encoding="utf-8")):
            offenders.append(rel)
    assert offenders == [], offenders
