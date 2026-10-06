"""작업 50: 회의록 5개 항목(AI 생성·직권 수정), 공통 변경 이력, 엑셀 내려받기, 마이그레이션.
네트워크 없음(Gemini 는 가짜 클라이언트). 임시 DB·업로드 폴더만 쓴다."""
import io
import json
from datetime import date

import pytest
from alembic import command
from openpyxl import load_workbook
from sqlalchemy import select, text

from app.db import make_engine
from app.models import ActionItem, Event, Meeting, MeetingMinutes
from app.models.minutes import MINUTES_FIELDS, NO_CONTENT
from app.pipeline import extractor as ex
from app.pipeline.extractor import ExtractionResult, GeminiExtractor, parse_minutes
from app.services import export_sheets as sheets
from app.services.processing import process_meeting
from tests.test_meeting_queries import BASE, make_meeting
from tests.test_upload_processing import _alembic, add_account, auth, env, upload  # noqa: F401  (env 는 픽스처)

NEW_VALUES = {"purpose": "새 목적", "discussion": "새 논의", "decisions": "새 결정", "risks": "새 리스크", "nextAgenda": "새 안건"}


# ---------------- 공용 ----------------
@pytest.fixture
def team(env):
    f = env["factory"]
    return {
        "lead": add_account(f, "고객사A", "lead", "manager", name="박총괄", ),
        "mgr2": add_account(f, "고객사A", "mgr2", "manager", name="정관리"),
        "exe": add_account(f, "고객사A", "boss", "executive", name="최임원"),
        "staff": add_account(f, "고객사A", "kim", "staff", name="김담당"),
        "outsider": add_account(f, "고객사B", "out", "executive", name="외부"),
    }


@pytest.fixture
def meeting_id(env, team):
    return make_meeting(
        env["factory"], team["lead"], participants=[team["staff"], team["mgr2"]], transcript="화자1: 안녕", status="confirmed",
        items=[
            {"title": "견적서 송부", "assignee_id": team["staff"].id, "due_date": date(2026, 10, 9), "status": "confirmed",
             "evidence_start_sec": 75.0, "evidence_quote": "견적서는 이번 주까지"},
            {"title": "설비 점검", "assignee_id": team["mgr2"].id, "due_undetermined": True},
            {"title": "삭제된 업무", "status": "deleted"},
            {"title": "=SUM(1,1)", "due_date": date(2026, 10, 10)},
        ],
    )


def call(env, method, account, path, **kwargs):
    return env["client"].request(method, path, headers=auth(account), **kwargs)


def minutes_row(env, mid) -> MeetingMinutes | None:
    with env["factory"]() as s:
        return s.get(MeetingMinutes, mid)


def history_events(env, mid) -> list[Event]:
    with env["factory"]() as s:
        return s.scalars(select(Event).where(Event.entity_id == mid, Event.event_type == "minutes.overridden")).all()


# ---------------- AI 생성과 "내용없음" ----------------
def test_processing_saves_five_items_with_no_content_for_empty(env, team):
    res = upload(env, team["lead"])
    assert process_meeting(res.json()["jobId"], session_factory=env["factory"]) == "completed"
    row = minutes_row(env, res.json()["meetingId"])
    assert row.purpose.startswith("STT 화자 분리") and row.decisions and row.next_agenda == "점심 메뉴 선정"
    assert row.risks == NO_CONTENT  # 가짜 응답이 비운 항목
    assert row.engine == "fake" and row.extract_model == "fake"
    detail = call(env, "GET", team["lead"], f"/api/meetings/{res.json()['meetingId']}").json()
    assert detail["minutes"]["risks"] == NO_CONTENT and detail["minutes"]["engine"] == "fake"
    assert detail["minutes"]["purpose"] == row.purpose and detail["minutes"]["updatedBy"] is None


def test_extractor_without_minutes_stores_no_content_and_keeps_items(env, team):
    class NoMinutes:
        def extract(self, transcript, held_at):
            fake = ex.FakeExtractor().extract(transcript, held_at)
            return ExtractionResult(items=fake.items, extract_model="m", prompt_version="p")

    res = upload(env, team["lead"])
    assert process_meeting(res.json()["jobId"], session_factory=env["factory"], extractor_factory=NoMinutes) == "completed"
    row = minutes_row(env, res.json()["meetingId"])
    assert all(getattr(row, key) == NO_CONTENT for key in MINUTES_FIELDS)
    with env["factory"]() as s:
        assert len(s.scalars(select(ActionItem).where(ActionItem.meeting_id == res.json()["meetingId"])).all()) == 2


class FakeGenai:
    """generate_content 를 흉내 내는 가짜 클라이언트. 프롬프트로 업무 추출/5개 항목을 구분해 응답한다."""

    def __init__(self, minutes_text: str | None, raise_on_minutes: bool = False):
        self.minutes_text, self.raise_on_minutes, self.calls = minutes_text, raise_on_minutes, []
        self.models = self

    def generate_content(self, *, model, contents, config):
        self.calls.append(contents)
        if "5개 항목" in contents:
            if self.raise_on_minutes:
                raise RuntimeError("boom")
            return type("R", (), {"text": self.minutes_text})()
        items = {"items": [{"task": "견적서 송부", "assignee": "김담당", "due_date": "2026-10-09",
                            "quote": {"speaker": "화자1", "timestamp": "00:01:15", "text": "견적서는 이번 주까지"}}]}
        return type("R", (), {"text": json.dumps(items)})()


@pytest.mark.parametrize("minutes_text, raises", [
    ("not json", False), ('{"purpose": 123, "risks": ["a"]}', False), (None, False), ("", True),
])
def test_bad_minutes_response_keeps_items_and_marks_no_content(env, team, minutes_text, raises):
    client = FakeGenai(minutes_text, raise_on_minutes=raises)
    res = upload(env, team["lead"])
    result = process_meeting(
        res.json()["jobId"], session_factory=env["factory"], extractor_factory=lambda: GeminiExtractor(client, "fake-model"),
    )
    assert result == "completed"
    assert len(client.calls) == 2  # 업무 추출 1회 + 5개 항목 1회
    row = minutes_row(env, res.json()["meetingId"])
    assert all(getattr(row, key) == NO_CONTENT for key in MINUTES_FIELDS)
    assert row.prompt_version == ""
    with env["factory"]() as s:
        items = s.scalars(select(ActionItem).where(ActionItem.meeting_id == res.json()["meetingId"])).all()
        assert [i.title for i in items] == ["견적서 송부"] and items[0].extract_model == "fake-model"


def test_good_gemini_minutes_response_and_partial_fields(env, team):
    client = FakeGenai(json.dumps({"purpose": " 목적 문장 ", "discussion": "", "decisions": "결정 A", "risks": None, "next_agenda": "다음"}))
    res = upload(env, team["lead"])
    process_meeting(res.json()["jobId"], session_factory=env["factory"], extractor_factory=lambda: GeminiExtractor(client, "m1"))
    row = minutes_row(env, res.json()["meetingId"])
    assert (row.purpose, row.discussion, row.decisions, row.risks, row.next_agenda) == ("목적 문장", NO_CONTENT, "결정 A", NO_CONTENT, "다음")
    assert row.extract_model == "m1" and row.prompt_version == ex.MINUTES_PROMPT_VERSION


def test_item_extraction_prompt_is_unchanged_by_minutes():
    # 업무 추출 규칙(프롬프트)은 5개 항목 때문에 바뀌지 않는다: 버전 해시가 고정이어야 한다
    assert "회의록의 5개 항목" not in ex.PROMPT_TEMPLATE
    assert parse_minutes('{"purpose": "a"}')["risks"] == NO_CONTENT


# ---------------- 직권 수정 ----------------
@pytest.mark.parametrize("who", ["lead", "exe"])
def test_override_allowed_for_lead_and_executive_with_history(env, team, meeting_id, who):
    res = call(env, "PATCH", team[who], f"/api/meetings/{meeting_id}/minutes", json={"purpose": "수정 목적", "risks": " 새 위험 "})
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["purpose"] == "수정 목적" and body["risks"] == "새 위험" and body["discussion"] == NO_CONTENT
    assert body["updatedBy"]["id"] == team[who].id and body["updatedAt"]
    [event] = history_events(env, meeting_id)
    assert event.actor_account_id == team[who].id
    assert event.payload == {"before": {"purpose": NO_CONTENT, "risks": NO_CONTENT}, "after": {"purpose": "수정 목적", "risks": "새 위험"}}
    # 상세 응답에도 반영(확정된 회의록이라 확정 이후 수정도 가능함을 함께 확인)
    assert call(env, "GET", team[who], f"/api/meetings/{meeting_id}").json()["minutes"]["purpose"] == "수정 목적"


def test_override_only_changed_fields_recorded_and_no_change_no_history(env, team, meeting_id):
    path = f"/api/meetings/{meeting_id}/minutes"
    call(env, "PATCH", team["lead"], path, json={"purpose": "A", "decisions": "B"})
    res = call(env, "PATCH", team["lead"], path, json={"purpose": "A", "decisions": "C"})  # purpose 는 그대로
    assert res.status_code == 200
    events = history_events(env, meeting_id)
    assert len(events) == 2 and events[1].payload == {"before": {"decisions": "B"}, "after": {"decisions": "C"}}
    assert call(env, "PATCH", team["lead"], path, json={"purpose": "A", "decisions": "C"}).status_code == 200
    assert len(history_events(env, meeting_id)) == 2  # 바뀐 것이 없으면 이력 없음
    # 빈 값은 "내용없음"으로 저장, 이미 내용없음인 항목을 비우는 것은 변경이 아님
    call(env, "PATCH", team["lead"], path, json={"purpose": "  ", "discussion": ""})
    assert minutes_row(env, meeting_id).purpose == NO_CONTENT
    assert len(history_events(env, meeting_id)) == 3


def test_override_denied_for_staff_other_manager_and_other_tenant(env, team, meeting_id):
    path = f"/api/meetings/{meeting_id}/minutes"
    assert call(env, "PATCH", team["staff"], path, json={"purpose": "x"}).status_code == 403
    assert call(env, "PATCH", team["mgr2"], path, json={"purpose": "x"}).status_code == 403  # 총괄 아닌 관리자
    assert call(env, "PATCH", team["outsider"], path, json={"purpose": "x"}).status_code == 404
    assert env["client"].patch(path, json={"purpose": "x"}).status_code == 401
    assert minutes_row(env, meeting_id) is None and history_events(env, meeting_id) == []


@pytest.mark.parametrize("body", [{}, {"purpose": None}, {"unknown": "x"}])
def test_override_requires_text_fields(env, team, meeting_id, body):
    assert call(env, "PATCH", team["lead"], f"/api/meetings/{meeting_id}/minutes", json=body).status_code == 400


def test_override_too_long_422(env, team, meeting_id):
    res = call(env, "PATCH", team["lead"], f"/api/meetings/{meeting_id}/minutes", json={"purpose": "가" * 5001})
    assert res.status_code == 422


def test_override_409_for_held_ended_deleted(env, team):
    for action, body in (("hold", {"reason": "검토"}), ("end", {"reason": "취소"}), ("delete", {"reason": "삭제"})):
        mid = make_meeting(env["factory"], team["lead"], transcript="x", status="confirmed", title=f"m-{action}")
        assert call(env, "POST", team["lead"], f"/api/meetings/{mid}/{action}", json=body).status_code == 200
        res = call(env, "PATCH", team["lead"], f"/api/meetings/{mid}/minutes", json={"purpose": "x"})
        assert res.status_code == 409, action
        assert history_events(env, mid) == []


# ---------------- 변경 이력 조회 ----------------
def test_history_listing_and_visibility(env, team, meeting_id):
    call(env, "PATCH", team["lead"], f"/api/meetings/{meeting_id}/minutes", json={"purpose": "P1"})
    call(env, "PATCH", team["exe"], f"/api/meetings/{meeting_id}/minutes", json={"purpose": "P2", "risks": "R"})
    res = call(env, "GET", team["staff"], f"/api/meetings/{meeting_id}/history")  # 담당자(참석자)도 열람 가능
    assert res.status_code == 200
    rows = res.json()
    assert [r["after"] for r in rows] == [{"purpose": "P2", "risks": "R"}, {"purpose": "P1"}]  # 최신순
    first = rows[0]
    assert set(first) == {"id", "targetType", "targetId", "kind", "before", "after", "changedBy", "changedAt"}
    assert first["kind"] == "직권 수정" and first["targetType"] == "meeting" and first["targetId"] == meeting_id
    assert first["changedBy"] == {"id": team["exe"].id, "name": "최임원"} and first["before"] == {"purpose": "P1", "risks": NO_CONTENT}
    assert call(env, "GET", team["outsider"], f"/api/meetings/{meeting_id}/history").status_code == 404
    assert env["client"].get(f"/api/meetings/{meeting_id}/history").status_code == 401


def test_history_hidden_from_staff_without_access_to_held_meeting(env, team, meeting_id):
    call(env, "POST", team["lead"], f"/api/meetings/{meeting_id}/hold", json={"reason": "검토"})
    assert call(env, "GET", team["staff"], f"/api/meetings/{meeting_id}/history").status_code == 404  # 담당자는 보류 회의록 열람 불가
    assert call(env, "GET", team["lead"], f"/api/meetings/{meeting_id}/history").status_code == 200


def test_history_structure_is_shared_events_table():
    # 별도 이력 표 없이 events 를 재사용한다
    from app.db import Base
    assert not [t for t in Base.metadata.tables if "history" in t or "change" in t]


# ---------------- 엑셀 ----------------
def download(env, account, mid):
    return call(env, "GET", account, f"/api/meetings/{mid}/export")


def test_export_workbook_sheets_columns_and_rows(env, team, meeting_id):
    call(env, "PATCH", team["lead"], f"/api/meetings/{meeting_id}/minutes", json=NEW_VALUES)
    res = download(env, team["staff"], meeting_id)  # 열람 가능한 담당자도 받을 수 있다
    assert res.status_code == 200 and res.headers["content-type"] == sheets.XLSX_MIME
    book = load_workbook(io.BytesIO(res.content))
    assert book.sheetnames == ["회의 정보", "업무"] == [sheets.SHEET_INFO, sheets.SHEET_TASKS]

    info = {row[0].value: row[1].value for row in book["회의 정보"].iter_rows(min_row=2)}
    assert [c.value for c in book["회의 정보"][1]] == ["항목", "내용"]
    assert list(info) == ["제목", "일시", "상태", "등록자", "확정 정보", "참석자", "목적", "주요 논의사항", "결정사항", "리스크", "다음 안건"]
    assert info["제목"] == "회의 0" and info["상태"] == "확정" and info["일시"] == sheets.local_text(BASE)
    assert info["등록자"] == "박총괄(lead)"
    assert "mgr2" in info["참석자"] and "kim" in info["참석자"]  # 이름과 계정 ID
    assert (info["목적"], info["주요 논의사항"], info["결정사항"], info["리스크"], info["다음 안건"]) == (
        "새 목적", "새 논의", "새 결정", "새 리스크", "새 안건")

    sheet = book["업무"]
    assert [c.value for c in sheet[1]] == [
        "업무 ID", "업무명", "담당자", "담당자 계정 ID", "기한", "상태", "근거 타임스탬프", "근거 인용문"]
    rows = [[("" if c.value is None else c.value) for c in r] for r in sheet.iter_rows(min_row=2)]  # 빈 칸은 읽을 때 None
    titles = [r[1] for r in rows]
    assert titles == ["견적서 송부", "설비 점검", "=SUM(1,1)"]  # 삭제된 업무 제외
    assert rows[0][2:] == ["김담당", "kim", "2026-10-09", "확정", "00:01:15", "견적서는 이번 주까지"]
    assert rows[1][2:5] == ["정관리", "mgr2", "미확정"] and rows[1][6] == ""
    assert rows[2][2] == "" and rows[2][3] == ""  # 담당자 없음
    # 수식처럼 보이는 내용도 글자로 저장된다(실행되지 않음)
    assert sheet.cell(row=4, column=2).data_type == "s"


def test_export_content_disposition_utf8_filename(env, team, meeting_id):
    res = download(env, team["lead"], meeting_id)
    disposition = res.headers["content-disposition"]
    assert disposition.startswith("attachment;") and f'filename="meeting-{meeting_id}.xlsx"' in disposition
    assert "filename*=UTF-8''" in disposition
    from urllib.parse import unquote
    assert unquote(disposition.split("filename*=UTF-8''")[1]) == f"회의 0_{sheets.local_text(BASE)[:10].replace('-', '')}.xlsx"
    disposition.encode("ascii")  # 헤더는 ASCII 로만 구성된다


def test_export_confirm_info_and_default_no_content(env, team):
    mid = make_meeting(env["factory"], team["lead"], transcript="x", status="awaiting_confirmation")
    book = load_workbook(io.BytesIO(download(env, team["lead"], mid).content))
    info = {row[0].value: row[1].value for row in book["회의 정보"].iter_rows(min_row=2)}
    assert info["확정 정보"] in ("", None) and info["상태"] == "확정 대기" and info["목적"] == NO_CONTENT
    assert book["업무"].max_row == 1  # 머리줄만


def test_export_denied_for_other_tenant_unrelated_staff_and_anonymous(env, team, meeting_id):
    assert download(env, team["outsider"], meeting_id).status_code == 404
    stranger = add_account(env["factory"], "고객사A", "lee", "staff", name="이무관")
    assert download(env, stranger, meeting_id).status_code == 404
    assert env["client"].get(f"/api/meetings/{meeting_id}/export").status_code == 401
    call(env, "POST", team["lead"], f"/api/meetings/{meeting_id}/hold", json={"reason": "검토"})
    assert download(env, team["staff"], meeting_id).status_code == 404  # 담당자는 보류 회의록 불가
    assert download(env, team["exe"], meeting_id).status_code == 200


# ---------------- 마이그레이션 ----------------
def _tables(url: str) -> set[str]:
    engine = make_engine(url)
    try:
        with engine.connect() as conn:
            return {r[0] for r in conn.execute(text("SELECT name FROM sqlite_master WHERE type='table'"))}
    finally:
        engine.dispose()


def test_migration_up_down_and_refuses_downgrade_with_rows(tmp_path):
    url = f"sqlite:///{(tmp_path / 'mig.db').as_posix()}"
    command.upgrade(_alembic(url), "head")
    assert "meeting_minutes" in _tables(url)
    command.downgrade(_alembic(url), "-1")  # 비어 있으면 되돌려진다
    assert "meeting_minutes" not in _tables(url)
    command.upgrade(_alembic(url), "head")

    # 행이 있으면 거부: 실제 회의록 흐름으로 행을 만든다
    from sqlalchemy.orm import sessionmaker
    engine = make_engine(url)
    factory = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)
    lead = add_account(factory, "고객사A", "lead", "manager")
    mid = make_meeting(factory, lead, transcript="x")
    with factory() as s:
        s.add(MeetingMinutes(meeting_id=mid, tenant_id=lead.tenant_id))
        s.commit()
    engine.dispose()
    with pytest.raises(RuntimeError, match="downgrade 거부"):
        command.downgrade(_alembic(url), "-1")
    assert "meeting_minutes" in _tables(url)
