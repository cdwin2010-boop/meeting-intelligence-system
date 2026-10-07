"""작업 51: 수정 회의록 업로드 갱신(미리보기·적용)과 수기 업무 등록. 네트워크 없음, 임시 DB·업로드 폴더만 쓴다.
업로드 파일은 작업 50 다운로드 API 로 받은 엑셀을 고쳐서 만든다(같은 형식을 읽는지 함께 확인)."""
import io
import json
from datetime import date

import pytest
from alembic import command
from openpyxl import load_workbook
from sqlalchemy import func, inspect, select, text
from sqlalchemy.orm import sessionmaker

from app.config import settings
from app.db import make_engine
from app.models import (
    ActionItem, Event, Meeting, MeetingGuestParticipant, MeetingMinutes, MeetingParticipant,
)
from app.services import export_sheets as sheets
from tests.test_meeting_queries import make_meeting
from tests.test_upload_processing import _alembic, add_account, auth, env  # noqa: F401  (env 는 픽스처)

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DUE = date(2026, 10, 9)


@pytest.fixture
def team(env):
    f = env["factory"]
    return {
        "lead": add_account(f, "고객사A", "lead", "manager", name="박총괄"),
        "mgr2": add_account(f, "고객사A", "mgr2", "manager", name="정관리"),
        "exe": add_account(f, "고객사A", "boss", "executive", name="최임원"),
        "kim": add_account(f, "고객사A", "kim", "staff", name="김담당"),
        "lee": add_account(f, "고객사A", "lee", "staff", name="이서연"),
        "dm1": add_account(f, "고객사A", "dm1", "staff", name="박동명"),
        "dm2": add_account(f, "고객사A", "dm2", "staff", name="박동명"),
        "outsider": add_account(f, "고객사B", "out", "executive", name="외부"),
    }


@pytest.fixture
def mid(env, team):
    """총괄 등록, 참석: kim·mgr2. 업무: 1 확정 대기(김담당) 2 확정 대기(담당자·기한 없음) 3 확정 4 종결 5 삭제"""
    return make_meeting(
        env["factory"], team["lead"], participants=[team["kim"], team["mgr2"]], transcript="화자1: 안녕", status="confirmed",
        items=[
            {"title": "견적서 송부", "assignee_id": team["kim"].id, "due_date": DUE},
            {"title": "설비 점검"},
            {"title": "확정 업무", "assignee_id": team["kim"].id, "due_date": DUE, "status": "confirmed", "confirm_kind": "manager"},
            {"title": "종결 업무", "assignee_id": team["kim"].id, "due_date": DUE, "status": "closed"},
            {"title": "삭제 업무", "status": "deleted"},
        ],
    )


def call(env, method, account, path, **kwargs):
    return env["client"].request(method, path, headers=auth(account) if account else {}, **kwargs)


def item_ids(env, mid) -> list[int]:
    with env["factory"]() as s:
        return list(s.scalars(select(ActionItem.id).where(ActionItem.meeting_id == mid).order_by(ActionItem.id)))


def get_item(env, item_id) -> ActionItem:
    with env["factory"]() as s:
        return s.get(ActionItem, item_id)


def exported(env, account, mid) -> bytes:
    res = call(env, "GET", account, f"/api/meetings/{mid}/export")
    assert res.status_code == 200
    return res.content


def edit(content: bytes, fn) -> bytes:
    """다운로드 엑셀을 고쳐 다시 저장한다. fn(info_sheet, task_sheet)"""
    book = load_workbook(io.BytesIO(content))
    fn(book[sheets.SHEET_INFO], book[sheets.SHEET_TASKS])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def task_row(sheet, item_id) -> int:
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row=row, column=1).value == item_id:
            return row
    raise AssertionError(f"업무 ID {item_id} 행 없음")


def col(name: str) -> int:
    return sheets.TASK_COLUMNS.index(name) + 1


def info_row(sheet, label: str) -> int:
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row=row, column=1).value == label:
            return row
    raise AssertionError(label)


def upload_call(env, account, mid, content, mode="preview", choices=None, name="수정.xlsx"):
    data = {"choices": json.dumps(choices)} if choices is not None else {}
    return call(env, "POST", account, f"/api/meetings/{mid}/update-upload/{mode}", data=data, files={"file": (name, content, XLSX)})


def counts(env) -> tuple[int, int, int, int]:
    with env["factory"]() as s:
        return tuple(s.scalar(select(func.count()).select_from(m)) for m in (ActionItem, Event, MeetingParticipant, MeetingGuestParticipant))


def events(env, kind) -> list[Event]:
    with env["factory"]() as s:
        return s.scalars(select(Event).where(Event.event_type == kind).order_by(Event.id)).all()


# ---------------- 권한 ----------------
@pytest.mark.parametrize("mode", ["preview", "apply"])
def test_upload_permissions(env, team, mid, mode):
    content = exported(env, team["lead"], mid)
    for who in ("lead", "exe"):
        assert upload_call(env, team[who], mid, content, mode).status_code == 200, who
    assert upload_call(env, team["mgr2"], mid, content, mode).status_code == 403  # 총괄 아닌 관리자
    assert upload_call(env, team["kim"], mid, content, mode).status_code == 403  # 담당자
    assert upload_call(env, team["outsider"], mid, content, mode).status_code == 404  # 다른 회사
    assert upload_call(env, None, mid, content, mode).status_code == 401


def test_roundtrip_unchanged_file_changes_nothing(env, team, mid):
    res = upload_call(env, team["lead"], mid, exported(env, team["lead"], mid))
    body = res.json()
    assert body["canApply"] is True and body["errors"] == [] and body["items"]["updates"] == [] and body["items"]["added"] == []
    assert body["participants"] is None and body["minutes"] == {}
    assert sorted(body["items"]["unchanged"]) == [item_ids(env, mid)[0], item_ids(env, mid)[1]]
    assert sorted(s["itemId"] for s in body["items"]["skipped"]) == item_ids(env, mid)[2:4]  # 확정·종결(삭제 업무는 내려받기에 없음)


# ---------------- 업무 갱신 ----------------
def test_update_pending_item_title_assignee_due_with_history(env, team, mid):
    ids = item_ids(env, mid)

    def change(info, tasks):
        row = task_row(tasks, ids[0])
        tasks.cell(row=row, column=col(sheets.COL_TITLE), value="견적서 송부(수정)")
        tasks.cell(row=row, column=col(sheets.COL_ASSIGNEE), value="이서연")
        tasks.cell(row=row, column=col(sheets.COL_ASSIGNEE_LOGIN), value="lee")
        tasks.cell(row=row, column=col(sheets.COL_DUE), value="2026-12-01")
        row2 = task_row(tasks, ids[1])
        tasks.cell(row=row2, column=col(sheets.COL_DUE), value="미확정")

    content = edit(exported(env, team["lead"], mid), change)
    before = counts(env)
    preview = upload_call(env, team["lead"], mid, content).json()
    assert counts(env) == before  # 미리보기는 아무것도 저장하지 않는다
    [update] = [u for u in preview["items"]["updates"] if u["itemId"] == ids[0]]
    assert update["before"] == {"title": "견적서 송부", "assigneeId": team["kim"].id, "dueDate": "2026-10-09"}
    assert update["after"] == {"title": "견적서 송부(수정)", "assigneeId": team["lee"].id, "dueDate": "2026-12-01"}

    res = upload_call(env, team["lead"], mid, content, "apply")
    assert res.status_code == 200, res.text
    body = res.json()
    assert sorted(body["updatedItemIds"]) == ids[:2] and body["addedItemIds"] == []
    one, two = get_item(env, ids[0]), get_item(env, ids[1])
    assert (one.title, one.assignee_id, one.due_date) == ("견적서 송부(수정)", team["lee"].id, date(2026, 12, 1))
    assert two.due_undetermined is True and two.due_date is None and two.assignee_id is None  # 담당자 칸이 비어 지우지 않음(원래 없음)
    recorded = events(env, "item.upload_updated")
    assert len(recorded) == 2 and {e.payload["batchId"] for e in recorded} == {body["batchId"]}  # 같은 업로드 묶음
    assert recorded[0].actor_account_id == team["lead"].id and recorded[0].payload["source"] == "upload"
    assert recorded[0].payload["after"]["title"] == "견적서 송부(수정)"
    history = call(env, "GET", team["lead"], f"/api/meetings/{mid}/history").json()
    assert {h["kindCode"] for h in history} == {"item.upload_updated"} and {h["kind"] for h in history} == {"업무 갱신"}
    assert history[0]["batchId"] == body["batchId"] and history[0]["targetType"] == "action_item"


def test_confirmed_closed_deleted_items_are_skipped_with_reason(env, team, mid):
    ids = item_ids(env, mid)

    def change(info, tasks):
        for item_id in ids[2:4]:
            tasks.cell(row=task_row(tasks, item_id), column=col(sheets.COL_TITLE), value="바꾸려는 이름")
        row = tasks.max_row + 1  # 삭제된 업무 ID 를 직접 적은 행
        tasks.cell(row=row, column=col(sheets.COL_ID), value=ids[4])
        tasks.cell(row=row, column=col(sheets.COL_TITLE), value="삭제된 업무 수정")

    content = edit(exported(env, team["lead"], mid), change)
    body = upload_call(env, team["lead"], mid, content, "apply").json()
    reasons = {s["itemId"]: s["reason"] for s in body["skipped"]}
    assert reasons == {ids[2]: "확정된 업무라 바꾸지 않았습니다", ids[3]: "종결된 업무라 바꾸지 않았습니다",
                       ids[4]: "삭제된 업무라 바꾸지 않았습니다"}
    assert [get_item(env, i).title for i in ids[2:]] == ["확정 업무", "종결 업무", "삭제 업무"]
    assert events(env, "item.upload_updated") == []


def test_new_row_is_added_at_bottom_as_manual_with_authority_evidence(env, team, mid):
    def change(info, tasks):
        row = tasks.max_row + 1
        tasks.cell(row=row, column=col(sheets.COL_TITLE), value="신규 업무")
        tasks.cell(row=row, column=col(sheets.COL_ASSIGNEE), value="이서연")  # 이름만(계정 ID 없음)
        tasks.cell(row=row, column=col(sheets.COL_DUE), value="2026-11-20")
        tasks.cell(row=row + 1, column=col(sheets.COL_TITLE), value="빈 칸 신규")  # 담당자·기한 비움 -> 보완 필요

    before_ids = item_ids(env, mid)
    body = upload_call(env, team["lead"], mid, edit(exported(env, team["lead"], mid), change), "apply").json()
    new_ids = body["addedItemIds"]
    assert len(new_ids) == 2 and item_ids(env, mid) == before_ids + new_ids  # 목록 맨 아래
    first, second = get_item(env, new_ids[0]), get_item(env, new_ids[1])
    assert (first.title, first.assignee_id, first.due_date, first.status) == ("신규 업무", team["lee"].id, date(2026, 11, 20), "pending")
    assert first.evidence_quote == "등록자 직권 지정" and first.extract_model == "manual" and first.evidence_start_sec is None
    assert second.assignee_id is None and second.due_date is None and second.needs_supplement  # 빈 칸은 보완 필요
    added = events(env, "item.manual_added")
    assert [e.entity_id for e in added] == new_ids and added[0].actor_account_id == team["lead"].id
    assert added[0].payload["source"] == "upload" and added[0].payload["batchId"] == body["batchId"]
    detail = call(env, "GET", team["lead"], f"/api/meetings/{mid}").json()
    shown = {i["id"]: i for i in detail["actionItems"]}
    assert shown[new_ids[0]]["origin"] == "manual" and shown[new_ids[0]]["evidenceQuote"] == "등록자 직권 지정"
    assert shown[before_ids[0]]["origin"] == "ai"


def test_items_missing_from_file_are_kept_and_blank_cells_do_not_erase(env, team, mid):
    ids = item_ids(env, mid)

    def change(info, tasks):
        tasks.delete_rows(task_row(tasks, ids[1]))  # 파일에서 업무 2 행 삭제
        row = task_row(tasks, ids[0])
        for name in (sheets.COL_TITLE, sheets.COL_ASSIGNEE, sheets.COL_ASSIGNEE_LOGIN, sheets.COL_DUE):
            tasks.cell(row=row, column=col(name), value=None)

    body = upload_call(env, team["lead"], mid, edit(exported(env, team["lead"], mid), change), "apply").json()
    assert body["updatedItemIds"] == [] and ids[0] in body["unchangedItemIds"]
    one = get_item(env, ids[0])
    assert (one.title, one.assignee_id, one.due_date) == ("견적서 송부", team["kim"].id, DUE)
    assert get_item(env, ids[1]).status == "pending" and ids[1] in item_ids(env, mid)  # 파일에 없어도 삭제하지 않는다


def test_status_and_evidence_columns_are_ignored(env, team, mid):
    ids = item_ids(env, mid)

    def change(info, tasks):
        row = task_row(tasks, ids[0])
        tasks.cell(row=row, column=col(sheets.COL_STATUS), value="종결")
        tasks.cell(row=row, column=col(sheets.COL_EVIDENCE_QUOTE), value="바뀐 근거")

    body = upload_call(env, team["lead"], mid, edit(exported(env, team["lead"], mid), change), "apply").json()
    assert body["updatedItemIds"] == [] and get_item(env, ids[0]).status == "pending" and get_item(env, ids[0]).evidence_quote is None


# ---------------- 담당자 찾기 ----------------
def test_assignee_resolution_login_name_unregistered_and_ambiguous(env, team, mid):
    ids = item_ids(env, mid)

    def change(info, tasks):
        row = task_row(tasks, ids[1])
        tasks.cell(row=row, column=col(sheets.COL_ASSIGNEE), value="이서연")  # 이름만, 계정 하나
        new = tasks.max_row + 1
        tasks.cell(row=new, column=col(sheets.COL_TITLE), value="미등록 담당")
        tasks.cell(row=new, column=col(sheets.COL_ASSIGNEE), value="없는사람")
        tasks.cell(row=new + 1, column=col(sheets.COL_TITLE), value="ID 우선")
        tasks.cell(row=new + 1, column=col(sheets.COL_ASSIGNEE), value="엉뚱한 이름")
        tasks.cell(row=new + 1, column=col(sheets.COL_ASSIGNEE_LOGIN), value="kim")
        tasks.cell(row=new + 2, column=col(sheets.COL_TITLE), value="외부 계정")
        tasks.cell(row=new + 2, column=col(sheets.COL_ASSIGNEE_LOGIN), value="out")  # 다른 회사 계정은 담당자가 될 수 없다

    content = edit(exported(env, team["lead"], mid), change)
    preview = upload_call(env, team["lead"], mid, content).json()
    assert preview["canApply"] is True
    assert sum("담당자 계정을 찾을 수 없어" in w["message"] for w in preview["warnings"]) >= 1
    body = upload_call(env, team["lead"], mid, content, "apply").json()
    assert get_item(env, ids[1]).assignee_id == team["lee"].id
    added = [get_item(env, i) for i in body["addedItemIds"]]
    assert [a.assignee_id for a in added] == [None, team["kim"].id, None]  # 없는 이름·다른 회사 계정은 담당자 미지정(보완 필요)
    assert added[0].needs_supplement and "assignee" in added[0].missing_fields


def test_namesake_requires_choice_before_apply(env, team, mid):
    ids = item_ids(env, mid)

    def change(info, tasks):
        tasks.cell(row=task_row(tasks, ids[1]), column=col(sheets.COL_ASSIGNEE), value="박동명")

    content = edit(exported(env, team["lead"], mid), change)
    preview = upload_call(env, team["lead"], mid, content).json()
    assert preview["canApply"] is False and len(preview["ambiguities"]) == 1
    [amb] = preview["ambiguities"]
    assert amb["name"] == "박동명" and {c["loginId"] for c in amb["candidates"]} == {"dm1", "dm2"}
    before = counts(env)
    refused = upload_call(env, team["lead"], mid, content, "apply")
    assert refused.status_code == 400 and refused.json()["detail"]["ambiguities"] and counts(env) == before
    bad = upload_call(env, team["lead"], mid, content, "apply", choices={amb["key"]: team["kim"].id})  # 후보가 아닌 계정
    assert bad.status_code == 400 and counts(env) == before
    ok = upload_call(env, team["lead"], mid, content, "apply", choices={amb["key"]: team["dm2"].id})
    assert ok.status_code == 200 and get_item(env, ids[1]).assignee_id == team["dm2"].id
    # 선택값을 같이 보낸 미리보기는 동명이인이 해소된다
    assert upload_call(env, team["lead"], mid, content, "preview", choices={amb["key"]: team["dm1"].id}).json()["ambiguities"] == []


# ---------------- 참석자 ----------------
def test_participants_replaced_with_unregistered_names_and_view_impact_in_history(env, team):
    f = env["factory"]
    meeting = make_meeting(f, team["lead"], participants=[team["kim"], team["lee"]], transcript="x", status="confirmed",
                           items=[{"title": "김담당 업무", "assignee_id": team["kim"].id, "due_date": DUE}])

    def change(info, tasks):
        info.cell(row=info_row(info, sheets.INFO_PARTICIPANTS), column=2, value="김담당(kim), 홍길동, 박총괄(lead)")

    content = edit(exported(env, team["lead"], meeting), change)
    preview = upload_call(env, team["lead"], meeting, content).json()["participants"]
    assert sorted(preview["added"]) == ["박총괄(lead)", "홍길동(미등록)"]
    assert [(r["name"], r["viewImpact"]) for r in preview["removed"]] == [("이서연", "열람 불가(참석자에서 빠져 볼 수 없게 됨)")]
    assert call(env, "GET", team["lee"], f"/api/meetings/{meeting}").status_code == 200  # 아직 적용 전

    assert upload_call(env, team["lead"], meeting, content, "apply").json()["participantsChanged"] is True
    with f() as s:
        assert {p.account_id for p in s.scalars(select(MeetingParticipant).where(MeetingParticipant.meeting_id == meeting))} == {
            team["kim"].id, team["lead"].id}
        assert [g.name for g in s.scalars(select(MeetingGuestParticipant))] == ["홍길동"]
    assert call(env, "GET", team["lee"], f"/api/meetings/{meeting}").status_code == 404  # 빠진 사람은 볼 수 없다(미등록 이름은 권한 무관)
    [event] = events(env, "participants.overridden")
    assert event.payload["viewImpact"][0]["accountId"] == team["lee"].id and "열람 불가" in event.payload["viewImpact"][0]["viewImpact"]
    assert "이서연(lee)" in event.payload["before"]["participants"][1] or "이서연(lee)" in event.payload["before"]["participants"]
    assert "홍길동(미등록)" in event.payload["after"]["participants"] and event.actor_account_id == team["lead"].id
    detail = call(env, "GET", team["lead"], f"/api/meetings/{meeting}").json()
    assert detail["guestParticipants"] == ["홍길동"]
    # 다시 내려받으면 "이름(미등록)"로 나오고, 그 파일을 올리면 변경이 없다
    again = exported(env, team["lead"], meeting)
    info = {r[0].value: r[1].value for r in load_workbook(io.BytesIO(again))[sheets.SHEET_INFO].iter_rows(min_row=2)}
    assert "홍길동(미등록)" in info[sheets.INFO_PARTICIPANTS]
    assert upload_call(env, team["lead"], meeting, again).json()["participants"] is None


def test_participant_namesake_needs_choice(env, team, mid):
    def change(info, tasks):
        info.cell(row=info_row(info, sheets.INFO_PARTICIPANTS), column=2, value="박동명")

    content = edit(exported(env, team["lead"], mid), change)
    preview = upload_call(env, team["lead"], mid, content).json()
    [amb] = preview["ambiguities"]
    assert amb["key"] == "participant:박동명" and upload_call(env, team["lead"], mid, content, "apply").status_code == 400
    assert upload_call(env, team["lead"], mid, content, "apply", choices={amb["key"]: team["dm1"].id}).status_code == 200
    with env["factory"]() as s:
        assert [p.account_id for p in s.scalars(select(MeetingParticipant).where(MeetingParticipant.meeting_id == mid))] == [team["dm1"].id]


# ---------------- 5개 항목 ----------------
def test_minutes_changed_by_upload_after_confirmation_with_history(env, team, mid):
    def change(info, tasks):
        info.cell(row=info_row(info, "목적"), column=2, value="업로드로 바꾼 목적")
        info.cell(row=info_row(info, "리스크"), column=2, value="새 리스크")

    content = edit(exported(env, team["lead"], mid), change)  # 확정된 회의록
    assert upload_call(env, team["lead"], mid, content).json()["minutes"]["purpose"] == {"before": "내용없음", "after": "업로드로 바꾼 목적"}
    body = upload_call(env, team["exe"], mid, content, "apply").json()
    assert sorted(body["minutesChanged"]) == ["purpose", "risks"]
    with env["factory"]() as s:
        row = s.get(MeetingMinutes, mid)
        assert (row.purpose, row.risks, row.decisions) == ("업로드로 바꾼 목적", "새 리스크", "내용없음")
    [event] = events(env, "minutes.overridden")
    assert event.actor_account_id == team["exe"].id and event.payload["batchId"] == body["batchId"]
    assert event.payload["before"] == {"purpose": "내용없음", "risks": "내용없음"}
    history = call(env, "GET", team["lead"], f"/api/meetings/{mid}/history").json()
    assert history[0]["kind"] == "직권 수정" and history[0]["targetType"] == "meeting"


# ---------------- 오류와 전체 미적용 ----------------
def test_file_error_applies_nothing_and_lists_errors(env, team, mid):
    other = make_meeting(env["factory"], team["lead"], transcript="x", items=[{"title": "다른 회의 업무"}], title="다른 회의")
    other_item = item_ids(env, other)[0]
    ids = item_ids(env, mid)

    def change(info, tasks):
        tasks.cell(row=task_row(tasks, ids[0]), column=col(sheets.COL_TITLE), value="이건 적용되면 안 됨")  # 정상 행
        info.cell(row=info_row(info, "목적"), column=2, value="이것도 적용되면 안 됨")
        row = tasks.max_row + 1
        tasks.cell(row=row, column=col(sheets.COL_ID), value=other_item)  # 다른 회의록 업무
        tasks.cell(row=row + 1, column=col(sheets.COL_ID), value=999999)  # 없는 업무
        tasks.cell(row=row + 2, column=col(sheets.COL_ID), value="abc")
        tasks.cell(row=row + 3, column=col(sheets.COL_TITLE), value="기한 오류")
        tasks.cell(row=row + 3, column=col(sheets.COL_DUE), value="내일")

    content = edit(exported(env, team["lead"], mid), change)
    preview = upload_call(env, team["lead"], mid, content).json()
    assert preview["canApply"] is False and len(preview["errors"]) == 4
    before = counts(env)
    res = upload_call(env, team["lead"], mid, content, "apply")
    assert res.status_code == 400 and len(res.json()["detail"]["errors"]) == 4
    assert {e["sheet"] for e in res.json()["detail"]["errors"]} == {sheets.SHEET_TASKS}
    assert counts(env) == before and get_item(env, ids[0]).title == "견적서 송부" and get_item(env, other_item).title == "다른 회의 업무"
    with env["factory"]() as s:
        assert s.get(MeetingMinutes, mid) is None
    assert events(env, "item.upload_updated") == [] and events(env, "minutes.overridden") == []


def test_formula_cells_are_ignored_not_evaluated(env, team, mid):
    ids = item_ids(env, mid)

    def change(info, tasks):
        tasks.cell(row=task_row(tasks, ids[0]), column=col(sheets.COL_TITLE), value="=1+1")  # 수식 셀
        tasks.cell(row=task_row(tasks, ids[1]), column=col(sheets.COL_TITLE), value="'=글자")  # 글자(작은따옴표 포함)

    preview = upload_call(env, team["lead"], mid, edit(exported(env, team["lead"], mid), change)).json()
    assert any("수식 셀" in w["message"] for w in preview["warnings"])
    updates = {u["itemId"]: u for u in preview["items"]["updates"]}
    assert ids[0] not in updates  # 수식은 값으로 쓰지 않음(기존 업무명 유지)
    assert updates[ids[1]]["after"]["title"] == "'=글자"


def test_text_that_looks_like_formula_from_export_roundtrips_as_text(env, team):
    f = env["factory"]
    meeting = make_meeting(f, team["lead"], transcript="x", items=[{"title": "=SUM(1,1)"}], status="confirmed")
    body = upload_call(env, team["lead"], meeting, exported(env, team["lead"], meeting)).json()
    assert body["items"]["updates"] == [] and body["errors"] == [] and body["warnings"] == []


def test_size_row_and_format_limits(env, team, mid, monkeypatch):
    content = exported(env, team["lead"], mid)
    monkeypatch.setattr(settings, "update_max_file_kb", 1)
    assert upload_call(env, team["lead"], mid, content).status_code == 413
    monkeypatch.setattr(settings, "update_max_file_kb", 2048)
    monkeypatch.setattr(settings, "update_max_rows", 1)
    res = upload_call(env, team["lead"], mid, content, "apply")
    assert res.status_code == 400 and "너무 많습니다" in res.json()["detail"]
    monkeypatch.setattr(settings, "update_max_rows", 1000)
    assert upload_call(env, team["lead"], mid, b"not an excel").status_code == 400
    assert upload_call(env, team["lead"], mid, b"").status_code == 400


def test_missing_sheet_or_column_rejected(env, team, mid):
    content = exported(env, team["lead"], mid)

    def drop_sheet(info, tasks):
        info.parent.remove(tasks)

    def rename_column(info, tasks):
        tasks.cell(row=1, column=col(sheets.COL_TITLE), value="이름")

    def bad_info_header(info, tasks):
        info.cell(row=1, column=1, value="키")

    for fn, part in ((drop_sheet, "시트가 없습니다"), (rename_column, "열이 없습니다"), (bad_info_header, "머리줄")):
        res = upload_call(env, team["lead"], mid, edit(content, fn), "apply")
        assert res.status_code == 400 and part in res.json()["detail"], part


def test_choices_must_be_json_object(env, team, mid):
    res = call(env, "POST", team["lead"], f"/api/meetings/{mid}/update-upload/preview", data={"choices": "[1]"},
               files={"file": ("a.xlsx", exported(env, team["lead"], mid), XLSX)})
    assert res.status_code == 400


@pytest.mark.parametrize("action, body", [("hold", {"reason": "검토"}), ("end", {"reason": "취소"}), ("delete", {"reason": "삭제"})])
def test_locked_meetings_409(env, team, action, body):
    meeting = make_meeting(env["factory"], team["lead"], transcript="x", status="confirmed", title=f"m-{action}",
                           items=[{"title": "업무"}])
    content = exported(env, team["lead"], meeting)
    assert call(env, "POST", team["lead"], f"/api/meetings/{meeting}/{action}", json=body).status_code == 200
    for mode in ("preview", "apply"):
        assert upload_call(env, team["lead"], meeting, content, mode).status_code == 409
    res = call(env, "POST", team["lead"], f"/api/meetings/{meeting}/action-items", json={"title": "x"})
    assert res.status_code == 409


# ---------------- 수기 업무 등록 ----------------
def manual(env, account, mid, **body):
    return call(env, "POST", account, f"/api/meetings/{mid}/action-items", json=body)


def test_manual_item_permissions(env, team, mid):
    assert manual(env, team["lead"], mid, title="총괄 등록").status_code == 201
    assert manual(env, team["exe"], mid, title="지시자 등록").status_code == 201
    assert manual(env, team["mgr2"], mid, title="타 관리자").status_code == 403
    assert manual(env, team["kim"], mid, title="담당자").status_code == 403
    assert manual(env, team["outsider"], mid, title="타사").status_code == 404
    assert manual(env, None, mid, title="로그인 없음").status_code == 401
    assert len(item_ids(env, mid)) == 7


def test_manual_item_fields_origin_evidence_and_history(env, team, mid):
    res = manual(env, team["lead"], mid, title="  수기 업무  ", assigneeId=team["lee"].id, dueDate="2026-11-30")
    assert res.status_code == 201
    body = res.json()
    assert body["title"] == "수기 업무" and body["assignee"]["id"] == team["lee"].id and body["dueDate"] == "2026-11-30"
    assert body["origin"] == "manual" and body["evidenceQuote"] == "등록자 직권 지정" and body["evidenceStartSec"] is None
    assert body["status"] == "pending" and body["needsCompletion"] is False
    undetermined = manual(env, team["lead"], mid, title="미확정 기한", assigneeId=team["lee"].id, dueUndetermined=True).json()
    assert undetermined["dueUndetermined"] is True and undetermined["needsCompletion"] is False
    [first, second] = events(env, "item.manual_added")
    assert first.entity_id == body["id"] and first.actor_account_id == team["lead"].id and first.created_at is not None
    assert first.payload["source"] == "manual" and first.payload["after"]["title"] == "수기 업무"
    history = call(env, "GET", team["staff"] if "staff" in team else team["kim"], f"/api/meetings/{mid}/history").json()
    assert history[0]["kind"] == "직권 등록" and history[0]["changedBy"]["id"] == team["lead"].id
    # 목록 맨 아래, AI 추출 업무와 구분해서 표시
    detail = call(env, "GET", team["lead"], f"/api/meetings/{mid}").json()
    assert [i["origin"] for i in detail["actionItems"]][-2:] == ["manual", "manual"] and detail["actionItems"][0]["origin"] == "ai"


def test_manual_item_required_rules_same_as_ai_items(env, team, mid):
    assert manual(env, team["lead"], mid, title="").status_code == 422
    assert manual(env, team["lead"], mid, title="   ").status_code == 422
    assert manual(env, team["lead"], mid).status_code == 422
    assert manual(env, team["lead"], mid, title="가" * 501).status_code == 422
    assert manual(env, team["lead"], mid, title="x", assigneeId=team["outsider"].id).status_code == 400  # 다른 회사 계정
    assert manual(env, team["lead"], mid, title="x", assigneeId=999999).status_code == 400
    assert manual(env, team["lead"], mid, title="x", dueDate="2026-11-30", dueUndetermined=True).status_code == 400
    # 담당자·기한이 비면 AI 추출 업무처럼 보완 필요, 확정은 막힌다. 채우면 같은 절차로 확정된다
    incomplete = manual(env, team["lead"], mid, title="덜 채운 수기").json()
    assert incomplete["needsCompletion"] is True and incomplete["missingFields"] == ["assignee", "dueDate"]
    assert call(env, "POST", team["lead"], f"/api/action-items/{incomplete['id']}/confirm").status_code == 409
    fixed = call(env, "PATCH", team["lead"], f"/api/action-items/{incomplete['id']}", json={"assigneeId": team["lee"].id, "dueUndetermined": True})
    assert fixed.status_code == 200
    confirmed = call(env, "POST", team["lead"], f"/api/action-items/{incomplete['id']}/confirm")
    assert confirmed.status_code == 200 and confirmed.json()["status"] == "confirmed" and confirmed.json()["origin"] == "manual"


@pytest.mark.parametrize("status", ["failed", "no_content", "processing"])
def test_manual_item_works_when_transcription_or_extraction_failed(env, team, status):
    meeting = make_meeting(env["factory"], team["lead"], status=status, title=f"실패-{status}")  # 전사문·업무 없음
    res = manual(env, team["lead"], meeting, title="수기로 채움", assigneeId=team["kim"].id, dueUndetermined=True)
    assert res.status_code == 201 and res.json()["origin"] == "manual"
    assert call(env, "GET", team["lead"], f"/api/meetings/{meeting}").json()["actionItems"][0]["title"] == "수기로 채움"


def test_manual_item_exports_and_reuploads_as_existing_item(env, team, mid):
    created = manual(env, team["lead"], mid, title="수기", assigneeId=team["kim"].id, dueDate="2026-11-30").json()
    content = exported(env, team["lead"], mid)
    rows = {r[0].value: [c.value for c in r] for r in load_workbook(io.BytesIO(content))[sheets.SHEET_TASKS].iter_rows(min_row=2)}
    assert rows[created["id"]][-1] == "등록자 직권 지정"
    body = upload_call(env, team["lead"], mid, content).json()
    assert created["id"] in body["items"]["unchanged"] and body["items"]["added"] == []


# ---------------- 마이그레이션 ----------------
def _tables(url: str) -> set[str]:
    engine = make_engine(url)
    try:
        with engine.connect() as conn:
            return set(inspect(conn).get_table_names())
    finally:
        engine.dispose()


_BEFORE_GUESTS = "9e3b7c1d5a42"  # meeting_guest_participants 마이그레이션의 down_revision(뒤에 마이그레이션이 더 생겨도 이 표만 되돌린다)


def test_guest_participants_migration_up_down_and_refuses_with_rows(tmp_path):
    url = f"sqlite:///{(tmp_path / 'mig.db').as_posix()}"
    command.upgrade(_alembic(url), "head")
    assert "meeting_guest_participants" in _tables(url)
    command.downgrade(_alembic(url), _BEFORE_GUESTS)  # 비어 있으면 되돌려진다
    assert "meeting_guest_participants" not in _tables(url) and "meeting_minutes" in _tables(url)
    command.upgrade(_alembic(url), "head")

    engine = make_engine(url)
    factory = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)
    lead = add_account(factory, "고객사A", "lead", "manager")
    meeting = make_meeting(factory, lead, transcript="x")
    with factory() as s:
        s.add(MeetingGuestParticipant(tenant_id=lead.tenant_id, meeting_id=meeting, name="홍길동"))
        s.commit()
    engine.dispose()
    with pytest.raises(RuntimeError, match="downgrade 거부"):
        command.downgrade(_alembic(url), _BEFORE_GUESTS)
    assert "meeting_guest_participants" in _tables(url)
