"""자료 파일(엑셀 .xlsx) 등록(작업 84-1): 변환 규칙, 업로드 때 바로 거부, STT 건너뛰기·재처리, 응답 필드. 엑셀 파일은 테스트 코드에서 만든다."""
import datetime as dt
import io
import re
import zipfile

import pytest
from openpyxl import Workbook
from sqlalchemy import select

from app.config import settings
from app.models import ActionItem, Event, Meeting, SourceDocument, Transcript
from app.pipeline.extractor import FakeExtractor
from app.pipeline.fakes import FAKE_SCRIPT
from app.pipeline.transcript_text import TranscriptTextError, parse_transcript
from app.pipeline.transcript_xlsx import xlsx_to_text
from app.services.processing import process_meeting
from app.services.source_kind import TRANSCRIPT_TXT_ENGINE
from tests.test_transcript_txt import NoStt, post_txt
from tests.test_upload_processing import (  # noqa: F401  (env 는 픽스처)
    HELD_AT, BrokenExtractor, add_account, auth, env, saved_files,
)

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def workbook_bytes(rows, *, sheet="회의", extra=None, hidden_extra=None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for row in rows:
        ws.append(row)
    for name, data in (extra or {}).items():
        other = wb.create_sheet(name)
        for row in data:
            other.append(row)
    for name, data in (hidden_extra or {}).items():
        other = wb.create_sheet(name)
        for row in data:
            other.append(row)
        other.sheet_state = "hidden"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


HEADER = ["시각", "화자", "발언"]
SCRIPT_ROWS = [[ts, speaker, text] for ts, speaker, text in FAKE_SCRIPT]


def post_xlsx(env, account, content: bytes, *, filename="회의 자료.xlsx", source_kind="transcript_xlsx", title="엑셀 회의"):
    data = {"title": title, "heldAt": HELD_AT, "sourceKind": source_kind}
    return env["client"].post("/api/meetings/upload", headers=auth(account), data=data, files={"file": (filename, content, XLSX_TYPE)})


@pytest.fixture
def mgr(env):
    return add_account(env["factory"], "고객사A", "mgr", "manager", name="박관리")


def run(env, job_id, extractor=FakeExtractor):
    NoStt.called = 0
    return process_meeting(job_id, session_factory=env["factory"], stt_factory=NoStt(), extractor_factory=extractor)


# ---------------- 변환 규칙 ----------------
def test_header_columns_become_timed_lines_and_match_txt_result():
    text = xlsx_to_text(workbook_bytes([HEADER, *SCRIPT_ROWS]))
    lines = text.splitlines()
    assert lines[0] == f"[{FAKE_SCRIPT[0][0]}] {FAKE_SCRIPT[0][1]}: {FAKE_SCRIPT[0][2]}"
    from app.pipeline.fakes import fake_transcript

    assert parse_transcript(text).text == fake_transcript()  # 같은 내용의 txt 와 같은 전사문


def test_header_names_columns_in_any_order_and_english_names():
    rows = [["Text", "Speaker", "Time"], ["안녕하세요", "김대리", "00:01:05"], ["네 알겠습니다", "박팀장", "1:10"]]
    assert xlsx_to_text(workbook_bytes(rows)).splitlines() == ["[00:01:05] 김대리: 안녕하세요", "[00:01:10] 박팀장: 네 알겠습니다"]


def test_title_rows_above_header_are_skipped_and_blank_rows_ignored():
    rows = [["회의 기록"], [], ["일시", "2026-10-01"], HEADER, ["00:00:03", "김대리", "시작합니다"], [None, None, None], ["00:00:09", "박팀장", "네"]]
    assert xlsx_to_text(workbook_bytes(rows)).splitlines() == ["[00:00:03] 김대리: 시작합니다", "[00:00:09] 박팀장: 네"]


@pytest.mark.parametrize(
    "value, expected",
    [
        (dt.time(0, 1, 5), "[00:01:05]"),
        (dt.timedelta(minutes=2, seconds=3), "[00:02:03]"),
        (dt.datetime(2026, 10, 1, 1, 2, 3), "[01:02:03]"),
        (0.5, "[12:00:00]"),  # 엑셀 시간 값(하루 비율)
        (83, "[00:01:23]"),  # 초 단위 숫자
        ("12:34", "[00:12:34]"),
        ("[01:02:03]", "[01:02:03]"),
    ],
)
def test_time_cell_formats(value, expected):
    line = xlsx_to_text(workbook_bytes([HEADER, [value, "김대리", "발언"]]))
    assert line == f"{expected} 김대리: 발언"


def test_unreadable_time_gives_untimed_line_and_missing_speaker_carries_previous():
    rows = [HEADER, ["오전", "김대리", "첫 발언"], [None, None, "이어지는 발언"], ["00:00:20", "박팀장", "답변"]]
    assert xlsx_to_text(workbook_bytes(rows)).splitlines() == ["김대리: 첫 발언", "김대리: 이어지는 발언", "[00:00:20] 박팀장: 답변"]


def test_no_speaker_column_uses_default_speaker_and_keeps_rows_separate():
    rows = [["시간", "내용"], ["00:00:01", "첫째"], ["00:00:02", "둘째"]]
    assert xlsx_to_text(workbook_bytes(rows)).splitlines() == ["[00:00:01] 화자1: 첫째", "[00:00:02] 화자1: 둘째"]
    only_text = xlsx_to_text(workbook_bytes([["발언"], ["하나"], ["둘"]]))
    assert only_text.splitlines() == ["화자1: 하나", "화자1: 둘"]
    assert len(parse_transcript(only_text).segments) == 2


def test_multiple_visible_sheets_in_order_and_hidden_sheet_ignored():
    raw = workbook_bytes(
        [HEADER, ["00:00:01", "가", "첫 시트"]],
        extra={"둘째": [HEADER, ["00:00:05", "나", "둘째 시트"]]},
        hidden_extra={"숨김": [HEADER, ["00:00:09", "다", "숨긴 시트"]]},
    )
    assert xlsx_to_text(raw).splitlines() == ["[00:00:01] 가: 첫 시트", "[00:00:05] 나: 둘째 시트"]


def test_multiline_cell_and_speaker_punctuation_are_normalized():
    rows = [HEADER, ["00:00:01", "김 대리(팀)", "첫 줄\n둘째 줄:  끝"]]
    assert xlsx_to_text(workbook_bytes(rows)) == "[00:00:01] 김대리팀: 첫 줄 둘째 줄: 끝"


def test_formula_cells_use_stored_values_only_and_never_the_formula_text():
    wb = Workbook()
    ws = wb.active
    ws.append(HEADER)
    ws.append(["00:00:01", "김대리", "=1+1"])  # 저장된 계산 값이 없는 수식: 읽지 않는다
    ws.append(["00:00:02", "김대리", "=HYPERLINK(\"http://example.com\",\"x\")"])
    ws.append(["00:00:03", "김대리", "진짜 발언"])
    buffer = io.BytesIO()
    wb.save(buffer)
    text = xlsx_to_text(buffer.getvalue())
    assert text == "[00:00:03] 김대리: 진짜 발언"
    assert "=" not in text and "HYPERLINK" not in text


# ---------------- 거부 규칙 ----------------
def test_rejects_unreadable_structures():
    cases = {
        "빈 시트": (workbook_bytes([]), "빈 파일"),
        "머리글만": (workbook_bytes([HEADER]), "발언이 없습니다"),
        "열 이름 없음": (workbook_bytes([["a", "b"], ["1", "2"]]), "열 이름을 찾지 못했습니다"),
        "zip 아님": (b"plain text, not a workbook", "엑셀(.xlsx) 파일이 아닙니다"),
        "빈 바이트": (b"", "빈 파일"),
    }
    for label, (raw, message) in cases.items():
        with pytest.raises(TranscriptTextError) as info:
            xlsx_to_text(raw)
        assert message in str(info.value), label


def test_rejects_password_protected_and_macro_files():
    encrypted = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 600  # 암호가 걸린 엑셀은 OLE 컨테이너
    with pytest.raises(TranscriptTextError, match="비밀번호"):
        xlsx_to_text(encrypted)
    base = workbook_bytes([HEADER, ["00:00:01", "김대리", "발언"]])
    macro = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(base)) as src, zipfile.ZipFile(macro, "w") as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr("xl/vbaProject.bin", b"fake")
    with pytest.raises(TranscriptTextError, match="매크로"):
        xlsx_to_text(macro.getvalue())
    macro_type = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(base)) as src, zipfile.ZipFile(macro_type, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"spreadsheetml.sheet.main+xml", b"ms-excel.sheet.macroEnabled.main+xml")
            dst.writestr(item, data)
    with pytest.raises(TranscriptTextError, match="매크로"):
        xlsx_to_text(macro_type.getvalue())


def test_rejects_corrupt_zip_oversize_and_too_many_chars(monkeypatch):
    good = workbook_bytes([HEADER, ["00:00:01", "김대리", "발언"]])
    with pytest.raises(TranscriptTextError, match="읽을 수 없습니다"):
        xlsx_to_text(good[: len(good) // 2])
    monkeypatch.setattr(settings, "transcript_upload_max_bytes", 1024)
    with pytest.raises(TranscriptTextError, match="너무 큽니다"):
        xlsx_to_text(good)
    monkeypatch.undo()
    monkeypatch.setattr(settings, "transcript_max_chars", 10)
    with pytest.raises(TranscriptTextError, match="글자 수"):
        xlsx_to_text(workbook_bytes([HEADER, ["00:00:01", "김대리", "열 글자를 넘기는 아주 긴 발언입니다"]]))


# ---------------- 업로드·처리 ----------------
def test_upload_xlsx_is_stored_and_processed_without_stt(env, mgr):
    res = post_xlsx(env, mgr, workbook_bytes([HEADER, *SCRIPT_ROWS]))
    assert res.status_code == 202
    [saved] = saved_files(env)
    assert saved.suffix == ".xlsx" and re.fullmatch(r"\d+_\d{10}_001", saved.stem)
    with env["factory"]() as s:
        meeting = s.get(Meeting, res.json()["meetingId"])
        assert s.get(SourceDocument, meeting.source_document_id).file_path.endswith(".xlsx")
        created = s.scalars(select(Event).where(Event.event_type == "meeting.created")).one()
        assert created.payload["sourceKind"] == "transcript_xlsx"
    assert run(env, res.json()["jobId"]) == "completed"
    assert NoStt.called == 0
    with env["factory"]() as s:
        transcript = s.scalar(select(Transcript))
        assert transcript.stt_provider == TRANSCRIPT_TXT_ENGINE
        assert transcript.full_text.splitlines()[0].startswith("[00:00:03] 김도현:")
        assert len(list(s.scalars(select(ActionItem)))) == 2  # 이후 추출은 음성·txt 와 같다


def test_detail_fields_and_audio_requests(env, mgr):
    res = post_xlsx(env, mgr, workbook_bytes([HEADER, *SCRIPT_ROWS]))
    meeting_id = res.json()["meetingId"]
    run(env, res.json()["jobId"])
    detail = env["client"].get(f"/api/meetings/{meeting_id}", headers=auth(mgr)).json()
    assert detail["sourceKind"] == "transcript_xlsx" and detail["hasAudio"] is False
    assert env["client"].get(f"/api/meetings/{meeting_id}/audio-url", headers=auth(mgr)).status_code == 404
    # 근거 시각은 엑셀의 시각 열에서 온다
    assert any(item["evidenceStartSec"] is not None for item in detail["actionItems"])


def test_reprocess_uses_stored_transcript_without_stt(env, mgr):
    res = post_xlsx(env, mgr, workbook_bytes([HEADER, *SCRIPT_ROWS]))
    meeting_id = res.json()["meetingId"]
    assert run(env, res.json()["jobId"], extractor=BrokenExtractor) == "failed"
    again = env["client"].post(f"/api/meetings/{meeting_id}/reprocess", headers=auth(mgr))
    assert again.status_code == 202
    assert run(env, again.json()["jobId"]) == "completed"
    assert NoStt.called == 0


def test_timeless_xlsx_keeps_evidence_time_empty(env, mgr):
    rows = [["화자", "발언"], ["이서연", "네, 제가 다음 주 금요일까지 정리해서 공유하겠습니다."]]
    res = post_xlsx(env, mgr, workbook_bytes(rows))
    assert run(env, res.json()["jobId"]) == "completed"
    with env["factory"]() as s:
        assert all(i.evidence_start_sec is None for i in s.scalars(select(ActionItem)))


@pytest.mark.parametrize(
    "filename, content, source_kind, message",
    [
        ("자료.txt", b"x", "transcript_xlsx", ".xlsx 파일만"),
        ("자료.xlsm", b"x", "transcript_xlsx", "매크로"),
        ("자료.xls", b"x", "transcript_xlsx", ".xlsx 파일만"),
        ("자료.xlsx", b"not a workbook", "transcript_xlsx", "엑셀(.xlsx) 파일이 아닙니다"),
        ("자료.xlsx", b"x", "transcript_txt", ".txt 파일만"),
        ("자료.xlsx", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 100, "transcript_xlsx", "비밀번호"),
    ],
)
def test_upload_rejections_are_422_and_nothing_is_saved(env, mgr, filename, content, source_kind, message):
    res = post_xlsx(env, mgr, content, filename=filename, source_kind=source_kind)
    assert res.status_code == 422 and message in res.json()["detail"]
    assert saved_files(env) == []
    with env["factory"]() as s:
        assert s.scalars(select(Meeting)).all() == []


def test_upload_size_limit_and_empty_sheet(env, mgr, monkeypatch):
    monkeypatch.setattr(settings, "transcript_upload_max_bytes", 2000)
    big = workbook_bytes([HEADER, *[["00:00:01", "김대리", "발언 " * 50] for _ in range(100)]])
    assert len(big) > 2000
    res = post_xlsx(env, mgr, big)
    assert res.status_code == 422 and "너무 큽니다" in res.json()["detail"]
    monkeypatch.undo()
    res = post_xlsx(env, mgr, workbook_bytes([]))
    assert res.status_code == 422 and "빈 파일" in res.json()["detail"]
    assert saved_files(env) == []


def test_audio_upload_and_txt_upload_are_unchanged(env, mgr):
    res = post_txt(env, mgr, "화자1: 안녕하세요 오늘 회의 시작합니다".encode("utf-8"))
    assert res.status_code == 202
    bad = env["client"].post(
        "/api/meetings/upload", headers=auth(mgr), data={"title": "x", "heldAt": HELD_AT},
        files={"file": ("자료.xlsx", workbook_bytes([HEADER, ["00:00:01", "a", "b"]]), XLSX_TYPE)},
    )
    assert bad.status_code == 415  # 음성으로 올린 xlsx 는 기존대로 허용 형식 오류
