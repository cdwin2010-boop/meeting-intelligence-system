"""수정 회의록 엑셀 업로드 갱신(작업 51). 작업 50 다운로드 파일(시트 "회의 정보"·"업무")과 같은 형식을 읽는다.
시트·열 이름은 app/services/export_sheets.py 상수만 쓴다.

두 단계, 서버에 상태를 저장하지 않는다:
  build_plan(...)  파일을 검증하고 변경 예정 목록을 만든다. 아무것도 저장하지 않는다(미리보기).
  apply_plan(...)  같은 파일·동명이인 선택값으로 만든 계획을 한 트랜잭션에서 적용한다(커밋은 호출부, 하나라도 실패하면 전부 취소).

갱신 규칙: 확정 대기 업무만 업무명·담당자·기한을 파일 내용으로 갱신(확정·종결·삭제 업무는 건너뜀), 업무 ID 가 빈 행은 새 업무(수기 등록과 같은 방식),
파일에 없는 업무는 건드리지 않음, 빈 셀은 기존 값을 지우지 않음, 상태·근거 열은 읽기 전용이라 무시. 수식 셀은 계산하지 않고 무시(경고).
회의록 5개 항목과 참석자는 확정 이후에도 바뀐다(변경 구분 "직권 수정")."""
import io
import re
import uuid
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.auth.access import VIEW_ALL_RANKS
from app.config import settings
from app.models import Account, ActionItem, Meeting, MeetingGuestParticipant, MeetingParticipant, SourceDocument
from app.models.action_item import MANUAL_EVIDENCE, MANUAL_ORIGIN
from app.models.item_conditions import not_deleted_item
from app.services.supersession import superseded_ids
from app.models.minutes import MINUTES_FIELDS
from app.services import export_sheets as sheets
from app.services.history import (
    KIND_ITEM_MANUAL_ADD, KIND_ITEM_UPLOAD_UPDATE, KIND_PARTICIPANTS_OVERRIDE, record_change,
)
from app.services.minutes_update import apply_minutes_changes, current_minutes

TITLE_MAX = 500
MINUTES_MAX = 5000
NAME_MAX = 100
# 압축을 푼 크기 상한 = 파일 크기 상한의 이 배수(압축 폭탄 방지)
_UNCOMPRESSED_FACTOR = 50
_EDITABLE_STATUS = "pending"
_SKIP_REASON = {
    "confirmed": "확정된 업무라 바꾸지 않았습니다",
    "closed": "종결된 업무라 바꾸지 않았습니다",
    "deleted": "삭제된 업무라 바꾸지 않았습니다",
}
_SUPERSEDED_REASON = "대체됨: 대체된 업무라 바꾸지 않았습니다"
_TOKEN_SPLIT = re.compile(r"[,\n;]+")
_PERSON_TOKEN = re.compile(r"^(.*?)\s*\((.+)\)\s*$")


class UploadFileError(Exception):
    """파일 자체를 읽을 수 없음(형식·용량·행 수). status: 400 | 413"""

    def __init__(self, message: str, status: int = 400):
        super().__init__(message)
        self.message, self.status = message, status


@dataclass
class Issue:
    sheet: str
    row: int | None
    message: str

    def out(self) -> dict[str, Any]:
        return {"sheet": self.sheet, "row": self.row, "message": self.message}


@dataclass
class Plan:
    errors: list[Issue] = field(default_factory=list)
    warnings: list[Issue] = field(default_factory=list)
    ambiguities: list[dict[str, Any]] = field(default_factory=list)
    # 업무
    updates: list[dict[str, Any]] = field(default_factory=list)  # {item, row, before, after}
    unchanged: list[int] = field(default_factory=list)
    skipped: list[dict[str, Any]] = field(default_factory=list)
    new_items: list[dict[str, Any]] = field(default_factory=list)  # {row, title, assignee_id, due_date, due_undetermined}
    # 참석자: None 이면 바뀌지 않음
    participants: dict[str, Any] | None = None
    # 5개 항목: {칸: {"before":..,"after":..}}
    minutes: dict[str, dict[str, str]] = field(default_factory=dict)

    @property
    def can_apply(self) -> bool:
        return not self.errors and not self.ambiguities

    def preview(self) -> dict[str, Any]:
        participants = None
        if self.participants is not None:
            p = self.participants
            participants = {
                "before": p["before_labels"], "after": p["after_labels"],
                "added": p["added_labels"], "removed": p["removed"],
            }
        return {
            "canApply": self.can_apply,
            "errors": [e.out() for e in self.errors],
            "warnings": [w.out() for w in self.warnings],
            "ambiguities": self.ambiguities,
            "items": {
                "updates": [
                    {"itemId": u["item"].id, "row": u["row"], "title": u["item"].title, "before": u["before"], "after": u["after"]}
                    for u in self.updates
                ],
                "unchanged": self.unchanged,
                "skipped": self.skipped,
                "added": [
                    {"row": n["row"], "title": n["title"], "assigneeId": n["assignee_id"],
                     "dueDate": n["due_date"].isoformat() if n["due_date"] else None, "dueUndetermined": n["due_undetermined"],
                     "evidence": MANUAL_EVIDENCE}
                    for n in self.new_items
                ],
            },
            "participants": participants,
            "minutes": self.minutes,
        }


# ---------------- 파일 읽기 ----------------
@dataclass
class ParsedFile:
    info: dict[str, tuple[int, str]]  # 항목 이름 -> (행, 내용)
    task_rows: list[tuple[int, dict[str, Any]]]  # (엑셀 행 번호, 열 이름 -> 값)
    formula_cells: list[tuple[str, str]]  # (시트, 셀 위치)


def _cell_value(cell, sheet_name: str, formulas: list[tuple[str, str]]) -> Any:
    """수식 셀은 계산하지 않고 무시(None)한다. 그 밖에는 값 그대로."""
    if cell.data_type == "f":
        formulas.append((sheet_name, cell.coordinate))
        return None
    return cell.value


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return str(value).strip()


def read_file(data: bytes) -> ParsedFile:
    limit = settings.update_max_file_kb * 1024
    if len(data) > limit:
        raise UploadFileError(f"파일이 너무 큽니다(최대 {settings.update_max_file_kb}KB).", 413)
    if not data:
        raise UploadFileError("빈 파일입니다.")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist()) > limit * _UNCOMPRESSED_FACTOR:
                raise UploadFileError("파일 내용이 너무 큽니다.", 413)
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except UploadFileError:
        raise
    except Exception as exc:  # noqa: BLE001 — 형식 오류는 종류와 상관없이 같은 안내
        raise UploadFileError("엑셀(.xlsx) 파일이 아니거나 읽을 수 없습니다.") from exc

    formulas: list[tuple[str, str]] = []
    missing = [name for name in (sheets.SHEET_INFO, sheets.SHEET_TASKS) if name not in book.sheetnames]
    if missing:
        raise UploadFileError(f"시트가 없습니다: {', '.join(missing)}")

    # "회의 정보": 머리줄 [항목, 내용]
    info_sheet = book[sheets.SHEET_INFO]
    info: dict[str, tuple[int, str]] = {}
    for number, row in enumerate(info_sheet.iter_rows(max_col=2), start=1):
        values = [_cell_value(c, sheets.SHEET_INFO, formulas) for c in row]
        values += [None] * (2 - len(values))
        if number == 1:
            if tuple(_text(v) for v in values) != sheets.INFO_HEADER:
                raise UploadFileError(f"\"{sheets.SHEET_INFO}\" 시트의 머리줄이 {list(sheets.INFO_HEADER)} 이어야 합니다.")
            continue
        label = _text(values[0])
        if label:
            info[label] = (number, _text(values[1]))

    # "업무": 머리줄의 열 이름으로 위치를 찾는다(순서가 바뀌어도 읽는다)
    task_sheet = book[sheets.SHEET_TASKS]
    columns: dict[str, int] = {}
    task_rows: list[tuple[int, dict[str, Any]]] = []
    required = (sheets.COL_ID, sheets.COL_TITLE, sheets.COL_ASSIGNEE, sheets.COL_ASSIGNEE_LOGIN, sheets.COL_DUE)
    for number, row in enumerate(task_sheet.iter_rows(), start=1):
        values = [_cell_value(c, sheets.SHEET_TASKS, formulas) for c in row]
        if number == 1:
            columns = {_text(v): i for i, v in enumerate(values) if _text(v)}
            lacking = [name for name in required if name not in columns]
            if lacking:
                raise UploadFileError(f"\"{sheets.SHEET_TASKS}\" 시트에 열이 없습니다: {', '.join(lacking)}")
            continue
        record = {name: (values[index] if index < len(values) else None) for name, index in columns.items() if name in required}
        if all(_text(v) == "" for v in record.values()):
            continue  # 빈 줄
        task_rows.append((number, record))
        if len(task_rows) > settings.update_max_rows:
            raise UploadFileError(f"업무 행이 너무 많습니다(최대 {settings.update_max_rows}행).")
    book.close()
    return ParsedFile(info=info, task_rows=task_rows, formula_cells=formulas)


# ---------------- 계정 찾기 ----------------
@dataclass
class Resolver:
    """계정 ID(login_id) 우선, 없으면 이름. 같은 고객사 활성 계정만. 이름이 둘 이상이면 동명이인(선택 필요)."""

    session: Session
    tenant_id: int
    choices: dict[str, int]
    plan: Plan

    def _by_login(self, login_id: str) -> Account | None:
        return self.session.scalar(select(Account).where(
            Account.tenant_id == self.tenant_id, Account.login_id == login_id, Account.is_active.is_(True)))

    def _by_name(self, name: str) -> list[Account]:
        return list(self.session.scalars(select(Account).where(
            Account.tenant_id == self.tenant_id, Account.name == name, Account.is_active.is_(True)).order_by(Account.id)))

    def resolve(self, key: str, login_id: str, name: str, *, sheet: str, row: int | None) -> tuple[str, Account | None]:
        """반환 (상태, 계정): resolved | unresolved(찾지 못함, 경고) | ambiguous(선택 필요) | blank"""
        if not login_id and not name:
            return "blank", None
        if login_id:
            found = self._by_login(login_id)
            if found is not None:
                return "resolved", found
            self.plan.warnings.append(Issue(sheet, row, f"계정 ID \"{login_id}\"를 찾을 수 없어 이름으로 찾습니다"))
            if not name:
                return "unresolved", None
        matches = self._by_name(name)
        if len(matches) == 1:
            return "resolved", matches[0]
        if len(matches) > 1:
            chosen_id = self.choices.get(key)
            if chosen_id is not None:
                chosen = next((m for m in matches if m.id == chosen_id), None)
                if chosen is None:
                    self.plan.errors.append(Issue(sheet, row, f"\"{name}\" 동명이인 선택값이 후보에 없습니다"))
                    return "unresolved", None
                return "resolved", chosen
            if not any(a["key"] == key for a in self.plan.ambiguities):
                self.plan.ambiguities.append({
                    "key": key, "name": name, "sheet": sheet, "row": row,
                    "candidates": [{"id": m.id, "name": m.name, "loginId": m.login_id} for m in matches],
                })
            return "ambiguous", None
        return "unresolved", None


def _parse_due(value: Any, sheet: str, row: int, plan: Plan) -> tuple[str, date | None]:
    """반환 (종류, 날짜): blank | undetermined | date | error"""
    if isinstance(value, datetime):
        return "date", value.date()
    if isinstance(value, date):
        return "date", value
    text = _text(value)
    if not text:
        return "blank", None
    if text == sheets.DUE_UNDETERMINED:
        return "undetermined", None
    try:
        return "date", date.fromisoformat(text)
    except ValueError:
        plan.errors.append(Issue(sheet, row, f"기한 \"{text}\" 은(는) 날짜(YYYY-MM-DD)나 \"{sheets.DUE_UNDETERMINED}\"이 아닙니다"))
        return "error", None


def _snapshot(title: str, assignee_id: int | None, due: date | None, undetermined: bool) -> dict[str, Any]:
    # 업무 PATCH 사건(item.updated)과 같은 키
    return {"title": title, "assigneeId": assignee_id, "dueDate": due.isoformat() if due else None, "dueUndetermined": undetermined}


# ---------------- 계획 ----------------
def build_plan(session: Session, meeting: Meeting, data: bytes, choices: dict[str, int]) -> Plan:
    """파일 → 변경 계획. DB 는 읽기만 한다(저장 없음). 파일 자체 오류는 UploadFileError."""
    parsed = read_file(data)
    plan = Plan()
    resolver = Resolver(session, meeting.tenant_id, choices, plan)
    for sheet_name, coordinate in parsed.formula_cells:
        plan.warnings.append(Issue(sheet_name, None, f"수식 셀 {coordinate} 은(는) 계산하지 않고 무시했습니다"))

    _plan_minutes(session, meeting, parsed, plan)
    _plan_items(session, meeting, parsed, resolver, plan)
    _plan_participants(session, meeting, parsed, resolver, plan)
    return plan


def _plan_minutes(session: Session, meeting: Meeting, parsed: ParsedFile, plan: Plan) -> None:
    current = current_minutes(session, meeting)
    for key, label in MINUTES_FIELDS.items():
        found = parsed.info.get(label)
        if found is None or not found[1]:
            continue  # 없거나 빈 셀은 기존 값을 지우지 않는다
        row, value = found
        if len(value) > MINUTES_MAX:
            plan.errors.append(Issue(sheets.SHEET_INFO, row, f"\"{label}\" 은(는) {MINUTES_MAX}자 이하여야 합니다"))
        elif value != current[key]:
            plan.minutes[key] = {"before": current[key], "after": value}


def _plan_items(session: Session, meeting: Meeting, parsed: ParsedFile, resolver: Resolver, plan: Plan) -> None:
    items = {i.id: i for i in session.scalars(select(ActionItem).where(ActionItem.meeting_id == meeting.id))}
    superseded = superseded_ids(session, list(items))  # 대체된 업무는 확정·종결·삭제 업무처럼 건너뛴다
    seen: set[int] = set()
    sheet = sheets.SHEET_TASKS
    for row, record in parsed.task_rows:
        id_text = _text(record.get(sheets.COL_ID))
        title = _text(record.get(sheets.COL_TITLE))
        assignee_name = _text(record.get(sheets.COL_ASSIGNEE))
        assignee_login = _text(record.get(sheets.COL_ASSIGNEE_LOGIN))
        if len(title) > TITLE_MAX:
            plan.errors.append(Issue(sheet, row, f"업무명은 {TITLE_MAX}자 이하여야 합니다"))
            continue

        if not id_text:
            # 새 업무(목록 맨 아래에 추가, 근거는 "등록자 직권 지정")
            state, account = resolver.resolve(f"row:{row}", assignee_login, assignee_name, sheet=sheet, row=row)
            if state == "unresolved":
                plan.warnings.append(Issue(sheet, row, "담당자 계정을 찾을 수 없어 담당자를 비워 둡니다(보완 필요)"))
            kind, due = _parse_due(record.get(sheets.COL_DUE), sheet, row, plan)
            if kind == "error":
                continue
            plan.new_items.append({
                "row": row, "title": title, "assignee_id": account.id if account else None,
                "due_date": due, "due_undetermined": kind == "undetermined", "pending_choice": state == "ambiguous",
            })
            continue

        if not id_text.isdigit():
            plan.errors.append(Issue(sheet, row, f"업무 ID \"{id_text}\" 이(가) 숫자가 아닙니다"))
            continue
        item_id = int(id_text)
        item = items.get(item_id)
        if item is None:
            # 다른 회의록·다른 고객사·없는 ID 를 구분하지 않는다
            plan.errors.append(Issue(sheet, row, f"업무 ID {item_id} 은(는) 이 회의록의 업무가 아닙니다"))
            continue
        if item_id in seen:
            plan.errors.append(Issue(sheet, row, f"업무 ID {item_id} 이(가) 두 번 나옵니다"))
            continue
        seen.add(item_id)
        if item_id in superseded or item.status != _EDITABLE_STATUS:
            plan.skipped.append({"itemId": item_id, "row": row, "title": item.title, "status": item.status,
                                 "reason": _SUPERSEDED_REASON if item_id in superseded else _SKIP_REASON.get(item.status, "바꿀 수 없는 상태입니다")})
            continue

        before = _snapshot(item.title, item.assignee_id, item.due_date, item.due_undetermined)
        after = dict(before)
        if title:
            after["title"] = title
        state, account = resolver.resolve(f"row:{row}", assignee_login, assignee_name, sheet=sheet, row=row)
        if state == "resolved":
            after["assigneeId"] = account.id
        elif state == "unresolved":
            plan.warnings.append(Issue(sheet, row, "담당자 계정을 찾을 수 없어 기존 담당자를 그대로 둡니다"))
        kind, due = _parse_due(record.get(sheets.COL_DUE), sheet, row, plan)
        if kind == "error":
            continue
        if kind == "date":
            after["dueDate"], after["dueUndetermined"] = due.isoformat(), False
        elif kind == "undetermined":
            after["dueDate"], after["dueUndetermined"] = None, True
        changed = {k for k in before if before[k] != after[k]}
        if changed:
            plan.updates.append({
                "item": item, "row": row,
                "before": {k: before[k] for k in changed}, "after": {k: after[k] for k in changed},
            })
        else:
            plan.unchanged.append(item_id)


def _participant_label(name: str, login_id: str | None) -> str:
    return f"{name}({login_id if login_id else sheets.UNREGISTERED})"


def _plan_participants(session: Session, meeting: Meeting, parsed: ParsedFile, resolver: Resolver, plan: Plan) -> None:
    found = parsed.info.get(sheets.INFO_PARTICIPANTS)
    if found is None or not found[1]:
        return  # 빈 셀은 기존 참석자를 지우지 않는다
    row, cell = found
    accounts: dict[int, Account] = {}
    guests: list[str] = []
    for token in (t.strip() for t in _TOKEN_SPLIT.split(cell)):
        if not token:
            continue
        match = _PERSON_TOKEN.match(token)
        name, tag = (match.group(1).strip(), match.group(2).strip()) if match else (token, "")
        login_id = "" if tag == sheets.UNREGISTERED else tag
        if not name and not login_id:
            continue
        if len(name) > NAME_MAX:
            plan.errors.append(Issue(sheets.SHEET_INFO, row, f"참석자 이름은 {NAME_MAX}자 이하여야 합니다"))
            continue
        state, account = resolver.resolve(f"participant:{name or login_id}", login_id, name, sheet=sheets.SHEET_INFO, row=row)
        if state == "resolved":
            accounts[account.id] = account
        elif state == "unresolved" and name:
            if name not in guests:
                guests.append(name)  # 계정이 없는 이름: "이름(미등록)" 글자만 남긴다
        elif state == "unresolved":
            plan.warnings.append(Issue(sheets.SHEET_INFO, row, f"참석자 계정 ID \"{login_id}\" 을(를) 찾을 수 없어 건너뜁니다"))

    current_accounts = {
        a.id: a for a in session.scalars(
            select(Account).join(MeetingParticipant, MeetingParticipant.account_id == Account.id)
            .where(MeetingParticipant.meeting_id == meeting.id).order_by(Account.id))
    }
    current_guests = list(session.scalars(
        select(MeetingGuestParticipant.name).where(MeetingGuestParticipant.meeting_id == meeting.id).order_by(MeetingGuestParticipant.id)))
    if plan.ambiguities and any(a["key"].startswith("participant:") for a in plan.ambiguities):
        # 동명이인 선택 전에는 참석자 변경을 확정할 수 없다(선택 후 다시 계산)
        plan.participants = None
        return
    if set(current_accounts) == set(accounts) and set(current_guests) == set(guests):
        return
    removed_accounts = [a for aid, a in current_accounts.items() if aid not in accounts]
    plan.participants = {
        "accounts": accounts, "guests": guests,
        "before_labels": [_participant_label(a.name, a.login_id) for a in current_accounts.values()] + [_participant_label(g, None) for g in current_guests],
        "after_labels": [_participant_label(a.name, a.login_id) for a in accounts.values()] + [_participant_label(g, None) for g in guests],
        "added_labels": [_participant_label(a.name, a.login_id) for aid, a in accounts.items() if aid not in current_accounts]
        + [_participant_label(g, None) for g in guests if g not in current_guests],
        "removed_accounts": removed_accounts,
        "removed": [],
    }
    _fill_view_impact(session, meeting, plan)


def _fill_view_impact(session: Session, meeting: Meeting, plan: Plan) -> None:
    """참석자에서 빠지는 사람이 이 회의록을 계속 볼 수 있는지(계획 적용 뒤 기준)."""
    registrant = session.scalar(select(SourceDocument.registered_by).where(SourceDocument.id == meeting.source_document_id))
    assigned: set[int] = set()
    planned = {u["item"].id: u["after"].get("assigneeId", u["item"].assignee_id) for u in plan.updates}
    for item in session.scalars(select(ActionItem).where(ActionItem.meeting_id == meeting.id, not_deleted_item())):
        owner = planned.get(item.id, item.assignee_id)
        if owner is not None:
            assigned.add(owner)
    assigned |= {n["assignee_id"] for n in plan.new_items if n["assignee_id"] is not None}
    removed = []
    for account in plan.participants["removed_accounts"]:
        if account.rank in VIEW_ALL_RANKS:
            result = "열람 유지(관리자 이상)"
        elif account.id == registrant:
            result = "열람 유지(등록자)"
        elif account.id in assigned:
            result = "열람 유지(담당 업무)"
        else:
            result = "열람 불가(참석자에서 빠져 볼 수 없게 됨)"
        removed.append({"accountId": account.id, "name": account.name, "loginId": account.login_id, "viewImpact": result})
    plan.participants["removed"] = removed


# ---------------- 수기 등록 ----------------
def create_manual_item(
    session: Session, meeting: Meeting, actor: Account, *, title: str, assignee_id: int | None,
    due_date: date | None, due_undetermined: bool, source: str = "manual", batch_id: str | None = None,
) -> ActionItem:
    """수기·업로드 새 행 업무 1건(확정 대기). 근거는 "등록자 직권 지정", 등록자·시각은 이력(행위자·시각)에 남는다.
    필수 항목 규칙·확정 절차는 AI 추출 업무와 같다(빈 칸이 있으면 보완 필요). 커밋은 호출부."""
    item = ActionItem(
        tenant_id=meeting.tenant_id, meeting_id=meeting.id, title=title, assignee_id=assignee_id,
        due_date=due_date, due_undetermined=due_undetermined, status="pending",
        evidence_start_sec=None, evidence_quote=MANUAL_EVIDENCE, extract_model=MANUAL_ORIGIN,
    )
    session.add(item)
    session.flush()
    extra: dict[str, Any] = {"source": source}
    if batch_id:
        extra["batchId"] = batch_id
    record_change(
        session, tenant_id=meeting.tenant_id, target_type="action_item", target_id=item.id, kind=KIND_ITEM_MANUAL_ADD,
        actor_id=actor.id, before={}, after=_snapshot(title, assignee_id, due_date, due_undetermined), extra=extra,
    )
    return item


# ---------------- 적용 ----------------
def apply_plan(session: Session, meeting: Meeting, actor: Account, plan: Plan) -> dict[str, Any]:
    """계획을 한 트랜잭션에서 적용한다(커밋은 호출부). plan.can_apply 가 아니면 호출하지 않는다."""
    batch_id = uuid.uuid4().hex
    extra = {"source": "upload", "batchId": batch_id}

    changed_minutes: list[str] = []
    if plan.minutes:
        result = apply_minutes_changes(session, meeting, actor, {k: v["after"] for k, v in plan.minutes.items()}, extra=extra)
        changed_minutes = list(result["after"])

    participants_changed = False
    if plan.participants is not None:
        p = plan.participants
        for row in session.scalars(select(MeetingParticipant).where(MeetingParticipant.meeting_id == meeting.id)):
            if row.account_id not in p["accounts"]:
                session.delete(row)
        have = set(session.scalars(select(MeetingParticipant.account_id).where(MeetingParticipant.meeting_id == meeting.id)))
        for account_id in p["accounts"]:
            if account_id not in have:
                session.add(MeetingParticipant(meeting_id=meeting.id, account_id=account_id))
        for row in session.scalars(select(MeetingGuestParticipant).where(MeetingGuestParticipant.meeting_id == meeting.id)):
            if row.name not in p["guests"]:
                session.delete(row)
        have_guests = set(session.scalars(select(MeetingGuestParticipant.name).where(MeetingGuestParticipant.meeting_id == meeting.id)))
        for name in p["guests"]:
            if name not in have_guests:
                session.add(MeetingGuestParticipant(tenant_id=meeting.tenant_id, meeting_id=meeting.id, name=name, created_by=actor.id))
        session.flush()
        record_change(
            session, tenant_id=meeting.tenant_id, target_type="meeting", target_id=meeting.id, kind=KIND_PARTICIPANTS_OVERRIDE,
            actor_id=actor.id, before={"participants": p["before_labels"]}, after={"participants": p["after_labels"]},
            extra={**extra, "viewImpact": p["removed"]},
        )
        participants_changed = True

    updated_ids: list[int] = []
    for update in plan.updates:
        item: ActionItem = update["item"]
        after = update["after"]
        if "title" in after:
            item.title = after["title"]
        if "assigneeId" in after:
            item.assignee_id = after["assigneeId"]
        if "dueDate" in after:
            item.due_date = date.fromisoformat(after["dueDate"]) if after["dueDate"] else None
        if "dueUndetermined" in after:
            item.due_undetermined = after["dueUndetermined"]
        record_change(
            session, tenant_id=meeting.tenant_id, target_type="action_item", target_id=item.id, kind=KIND_ITEM_UPLOAD_UPDATE,
            actor_id=actor.id, before=update["before"], after=after, extra=extra,
        )
        updated_ids.append(item.id)

    added_ids = [
        create_manual_item(
            session, meeting, actor, title=n["title"], assignee_id=n["assignee_id"], due_date=n["due_date"],
            due_undetermined=n["due_undetermined"], source="upload", batch_id=batch_id,
        ).id
        for n in plan.new_items
    ]
    session.flush()
    return {
        "batchId": batch_id, "updatedItemIds": updated_ids, "addedItemIds": added_ids,
        "skipped": plan.skipped, "unchangedItemIds": plan.unchanged,
        "participantsChanged": participants_changed, "minutesChanged": changed_minutes,
        "warnings": [w.out() for w in plan.warnings],
    }
