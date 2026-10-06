"""회의록 엑셀(.xlsx) 내보내기: 시트·열 이름 상수와 파일 만들기.
이후 업로드 갱신이 같은 파일을 읽으므로 시트 이름과 열 이름은 여기 한 곳에서만 정한다.
- "회의 정보" 시트: 머리줄 [항목, 내용] + 항목마다 한 줄(INFO_ROWS 순서)
- "업무" 시트: 머리줄 TASK_COLUMNS + 업무마다 한 줄(삭제된 업무 제외)
셀 값은 모두 글자로 저장한다(=, +, -, @ 로 시작하는 회의 내용이 수식으로 실행되지 않게)."""
import io
from datetime import date, datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook

from app.config import settings
from app.models.minutes import MINUTES_FIELDS

SHEET_INFO = "회의 정보"
SHEET_TASKS = "업무"

INFO_HEADER = ("항목", "내용")
INFO_TITLE = "제목"
INFO_HELD_AT = "일시"
INFO_STATUS = "상태"
INFO_REGISTRANT = "등록자"
INFO_CONFIRM = "확정 정보"
INFO_PARTICIPANTS = "참석자"
# 5개 항목의 표시 이름은 MINUTES_FIELDS 의 값(목적·주요 논의사항·결정사항·리스크·다음 안건)을 그대로 쓴다
INFO_ROWS = (INFO_TITLE, INFO_HELD_AT, INFO_STATUS, INFO_REGISTRANT, INFO_CONFIRM, INFO_PARTICIPANTS, *MINUTES_FIELDS.values())

COL_ID = "업무 ID"
COL_TITLE = "업무명"
COL_ASSIGNEE = "담당자"
COL_ASSIGNEE_LOGIN = "담당자 계정 ID"
COL_DUE = "기한"
COL_STATUS = "상태"
COL_EVIDENCE_TIME = "근거 타임스탬프"
COL_EVIDENCE_QUOTE = "근거 인용문"
TASK_COLUMNS = (COL_ID, COL_TITLE, COL_ASSIGNEE, COL_ASSIGNEE_LOGIN, COL_DUE, COL_STATUS, COL_EVIDENCE_TIME, COL_EVIDENCE_QUOTE)

DUE_UNDETERMINED = "미확정"
MEETING_STATUS_LABEL = {
    "processing": "처리 중", "awaiting_confirmation": "확정 대기", "confirmed": "확정", "failed": "처리 실패", "no_content": "내용 없음",
}
CONFIRM_KIND_LABEL = {
    "manager": "관리자 확정", "registration": "등록 시 확정", "period_elapsed": "기간 경과 자동 확정", "due_reached": "기한 도래 자동 확정",
}
ITEM_STATUS_LABEL = {"pending": "확정 대기", "confirmed": "확정", "closed": "종결"}

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def local_text(moment: datetime | None) -> str:
    """UTC 시각 -> 서비스 현지 시각 YYYY-MM-DD HH:mm"""
    if moment is None:
        return ""
    return moment.astimezone(ZoneInfo(settings.app_timezone)).strftime("%Y-%m-%d %H:%M")


def offset_text(sec: float | None) -> str:
    if sec is None:
        return ""
    total = int(sec)
    return f"{total // 3600:02d}:{total % 3600 // 60:02d}:{total % 60:02d}"


def due_text(due: date | None, undetermined: bool) -> str:
    if undetermined:
        return DUE_UNDETERMINED
    return due.isoformat() if due else ""


def _put(sheet, row: int, column: int, value) -> None:
    cell = sheet.cell(row=row, column=column)
    if isinstance(value, str):
        cell.value = value
        cell.data_type = "s"  # 수식으로 해석하지 않는다
    else:
        cell.value = value


def build_workbook(info: dict[str, str], tasks: list[dict[str, object]]) -> bytes:
    """info: INFO_ROWS 이름 -> 값, tasks: TASK_COLUMNS 이름 -> 값 인 딕셔너리 목록."""
    book = Workbook()
    sheet = book.active
    sheet.title = SHEET_INFO
    for column, name in enumerate(INFO_HEADER, start=1):
        _put(sheet, 1, column, name)
    for row, name in enumerate(INFO_ROWS, start=2):
        _put(sheet, row, 1, name)
        _put(sheet, row, 2, info.get(name, ""))
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 80

    task_sheet = book.create_sheet(SHEET_TASKS)
    for column, name in enumerate(TASK_COLUMNS, start=1):
        _put(task_sheet, 1, column, name)
    for row, task in enumerate(tasks, start=2):
        for column, name in enumerate(TASK_COLUMNS, start=1):
            _put(task_sheet, row, column, task.get(name, ""))

    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
