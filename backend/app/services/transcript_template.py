"""엑셀 자료 파일 양식(작업 85-1). 저장소에 바이너리를 두지 않고 요청 때 openpyxl 로 즉시 만든다.
시트1 "회의 발언": 머리글(시각·화자·발언)만 — 예시 행을 넣지 않아 실수로 그대로 등록되는 것을 막는다.
시트2 "작성 안내": 작성법과 예시. 읽기 규칙(app.pipeline.transcript_xlsx)은 첫 시트부터 보이는 시트를 모두 읽으므로,
안내 시트는 머리글('발언' 열)이 없어 읽기에서 건너뛴다(예시 표의 머리글은 '발언' 이 아닌 이름을 쓴다)."""
import io

TEMPLATE_FILENAME = "회의록_자료_양식.xlsx"
DATA_SHEET = "회의 발언"
GUIDE_SHEET = "작성 안내"
HEADERS = ("시각", "화자", "발언")

_GUIDE_ROWS = [
    ("작성 안내",),
    ("",),
    ("1. '회의 발언' 시트의 시각·화자·발언 열에 회의 내용을 한 행에 한 발언씩 적습니다.",),
    ("2. 발언 열은 필수입니다. 발언이 비어 있는 행은 읽지 않습니다.",),
    ("3. 시각 열은 비워도 됩니다. mm:ss(예: 05:30) 또는 hh:mm:ss(예: 01:05:30) 형식으로 적으면 근거 위치 보기에 쓰입니다.",),
    ("4. 화자 열은 비워도 됩니다. 비우면 화자 표시 없이 적은 그대로 보입니다.",),
    ("5. 매크로가 들어 있는 파일(.xlsm)과 암호가 걸린 파일은 올릴 수 없습니다. .xlsx 로 저장해 주세요.",),
    ("6. 이 안내 시트는 읽지 않으므로 지우지 않아도 됩니다.",),
    ("",),
    ("예시(이 표는 읽지 않습니다)",),
    ("시간 예시", "이름 예시", "내용 예시"),
    ("00:00:05", "김대리", "오늘 회의는 예산 검토부터 시작하겠습니다."),
    ("00:00:20", "박팀장", "수요일까지 견적서를 공유하겠습니다."),
]


def build_transcript_template() -> bytes:
    from openpyxl import Workbook  # noqa: PLC0415 — 양식을 만들 때만 불러온다
    from openpyxl.styles import Font

    workbook = Workbook()
    data = workbook.active
    data.title = DATA_SHEET
    data.append(list(HEADERS))
    for cell in data[1]:
        cell.font = Font(bold=True)
    for letter, width in (("A", 12), ("B", 14), ("C", 80)):
        data.column_dimensions[letter].width = width
    data.freeze_panes = "A2"
    data.cell(row=1, column=1).number_format = "@"

    guide = workbook.create_sheet(GUIDE_SHEET)
    for row in _GUIDE_ROWS:
        guide.append(list(row))
    guide["A1"].font = Font(bold=True)
    guide.column_dimensions["A"].width = 110
    guide.column_dimensions["B"].width = 14
    guide.column_dimensions["C"].width = 60

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
