"""자료 파일(엑셀 .xlsx) 전사문 읽기(작업 84-1). 엑셀의 시트·행·열을 "[hh:mm:ss] 화자: 발언" 줄로 합친 글자로 바꾼다.
그 뒤는 txt 자료 파일과 같다(transcript_text.parse_transcript 가 STT 결과와 같은 모양으로 바꾼다).

열 규칙(설계 문서 D-26 부록): 첫 행(머리글)의 이름으로 열을 찾는다. 위쪽 10행 안에서 "발언" 열이 있는 행을 머리글로 본다.
  - 시각 열(선택): 시각·시간·타임스탬프·time·timestamp …  "mm:ss"·"hh:mm:ss" 글자, 엑셀 시간 값, 초 단위 숫자를 읽는다. 읽을 수 없는 값은 시각 없음.
  - 화자 열(선택): 화자·발언자·이름·speaker …  없으면 "화자1" 로 둔다(화자 열이 있어도 빈 칸은 직전 화자).
  - 발언 열(필수): 발언·내용·대화·text …  발언이 빈 행은 건너뛴다.
  머리글을 찾지 못하면 오류(업로드 때 422). 시트는 보이는 시트를 순서대로 모두 읽는다(빈 시트는 건너뛰고, 전부 비면 오류).
안전: 수식은 저장된 계산 값만 읽는다(계산하지 않음, 값이 저장되지 않은 수식은 빈 칸). 외부 링크는 따라가지 않는다(openpyxl 은 읽지 않는다).
  비밀번호가 걸린 파일(OLE 컨테이너)·매크로 포함 파일(vbaProject·macroEnabled)·zip 이 아닌 파일·압축 해제 크기가 과한 파일은 거부한다.
전사 본문은 로그·오류 문구에 넣지 않는다.
"""
import datetime as dt
import io
import re
import zipfile

from app.config import settings
from app.pipeline.transcript_text import DEFAULT_SPEAKER, TranscriptTextError

_OLE_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # 암호가 걸린 엑셀은 zip 이 아니라 OLE 컨테이너로 저장된다
_ZIP_SIGNATURES = (b"PK\x03\x04", b"PK\x05\x06")
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_ROWS = 100_000
HEADER_SEARCH_ROWS = 10
_MAX_LABEL_LEN = 20

_TIME_NAMES = {"시각", "시간", "타임스탬프", "타임코드", "시작시각", "시작", "time", "timestamp", "timecode", "start"}
_SPEAKER_NAMES = {"화자", "화자명", "발언자", "말한사람", "참석자", "이름", "성명", "speaker", "name"}
_TEXT_NAMES = {"발언", "발언내용", "내용", "대화", "대화내용", "텍스트", "text", "content", "utterance", "message"}
_TIME_TEXT = re.compile(r"^\[?\s*(\d{1,3}):(\d{2})(?::(\d{2}))?\s*\]?$")
_LABEL_STRIP = re.compile(r"[\s.,!?;:：()\[\]]+")


def _norm(value: object) -> str:
    return re.sub(r"\s+", "", str(value or "")).lower()


def _hms(total: int) -> str:
    return f"[{total // 3600:02d}:{total % 3600 // 60:02d}:{total % 60:02d}]"


def _time_label(value: object) -> str | None:
    """시각 셀 → "[hh:mm:ss]". 읽을 수 없으면 None."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, dt.datetime):
        return _hms(value.hour * 3600 + value.minute * 60 + value.second)
    if isinstance(value, dt.time):
        return _hms(value.hour * 3600 + value.minute * 60 + value.second)
    if isinstance(value, dt.timedelta):
        return _hms(int(value.total_seconds()))
    if isinstance(value, (int, float)):
        if value < 0:
            return None
        # 0 이상 1 미만은 엑셀의 하루 비율(시간 값), 그 이상은 초 단위로 본다
        return _hms(int(round(value * 86400)) if value < 1 else int(value))
    match = _TIME_TEXT.match(str(value).strip())
    if not match:
        return None
    first, second, third = int(match.group(1)), int(match.group(2)), match.group(3)
    if third is None:  # mm:ss
        return _hms(first * 60 + second) if second < 60 else None
    return _hms(first * 3600 + second * 60 + int(third)) if second < 60 and int(third) < 60 else None


def _speaker_label(value: object) -> str:
    # 시각이 없는 줄에서도 발언 라벨로 읽히도록 공백·문장부호를 없애고 짧게 자른다
    label = _LABEL_STRIP.sub("", str(value or ""))
    return label[:_MAX_LABEL_LEN]


def _text_of(value: object) -> str:
    return re.sub(r"\s+", " ", str(value if value is not None else "")).strip()


def _check_container(raw: bytes) -> None:
    if raw[:8] == _OLE_SIGNATURE:
        raise TranscriptTextError("비밀번호가 걸린 엑셀 파일은 올릴 수 없습니다. 비밀번호를 해제한 뒤 올려 주세요.")
    if raw[:4] not in _ZIP_SIGNATURES:
        raise TranscriptTextError("엑셀(.xlsx) 파일이 아닙니다.")
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            names = archive.namelist()
            if any(name.lower().endswith("vbaproject.bin") for name in names):
                raise TranscriptTextError("매크로가 들어 있는 엑셀 파일은 올릴 수 없습니다. 매크로를 제거한 .xlsx 로 저장해 주세요.")
            if "xl/workbook.xml" not in names:
                raise TranscriptTextError("엑셀(.xlsx) 파일이 아닙니다.")
            content_types = archive.read("[Content_Types].xml").decode("utf-8", errors="ignore") if "[Content_Types].xml" in names else ""
            if "macroEnabled" in content_types:
                raise TranscriptTextError("매크로가 들어 있는 엑셀 파일은 올릴 수 없습니다. 매크로를 제거한 .xlsx 로 저장해 주세요.")
            if sum(info.file_size for info in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                raise TranscriptTextError("엑셀 파일의 내용이 너무 큽니다.")
    except zipfile.BadZipFile:
        raise TranscriptTextError("엑셀(.xlsx) 파일을 읽을 수 없습니다(손상되었거나 형식이 다릅니다).") from None


def _find_header(rows: list[tuple]) -> tuple[int, int | None, int | None, int] | None:
    """위쪽 행에서 머리글을 찾는다. (행 위치, 시각 열, 화자 열, 발언 열) 또는 None."""
    for index, row in enumerate(rows[:HEADER_SEARCH_ROWS]):
        names = [_norm(cell) for cell in row]
        text_col = next((i for i, n in enumerate(names) if n in _TEXT_NAMES), None)
        if text_col is None:
            continue
        time_col = next((i for i, n in enumerate(names) if n in _TIME_NAMES), None)
        speaker_col = next((i for i, n in enumerate(names) if n in _SPEAKER_NAMES and i != text_col), None)
        return index, time_col, speaker_col, text_col
    return None


def xlsx_to_text(raw: bytes) -> str:
    """엑셀 바이트 → 전사문 글자("[hh:mm:ss] 화자: 발언" 줄). 규칙 위반은 TranscriptTextError(사용자에게 보여도 되는 문구)."""
    if len(raw) > settings.transcript_upload_max_bytes:
        raise TranscriptTextError(f"엑셀 파일이 너무 큽니다(최대 {settings.transcript_upload_max_bytes // 1024}KB).")
    if not raw:
        raise TranscriptTextError("빈 파일입니다(내용이 없습니다).")
    _check_container(raw)
    try:
        from openpyxl import load_workbook  # noqa: PLC0415 — 엑셀을 읽을 때만 불러온다

        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except TranscriptTextError:
        raise
    except Exception:  # noqa: BLE001 — openpyxl 의 다양한 형식 오류를 한 문구로
        raise TranscriptTextError("엑셀(.xlsx) 파일을 읽을 수 없습니다(손상되었거나 형식이 다릅니다).") from None

    lines: list[str] = []
    saw_data_sheet = False
    header_missing = False
    try:
        for sheet in workbook.worksheets:
            if getattr(sheet, "sheet_state", "visible") != "visible":
                continue
            rows: list[tuple] = []
            for count, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if count > MAX_ROWS:
                    raise TranscriptTextError("엑셀 행이 너무 많습니다.")
                if any(cell is not None and str(cell).strip() != "" for cell in row):
                    rows.append(row)
                else:
                    rows.append(tuple())  # 빈 행도 위치는 유지하되 아래에서 건너뛴다
            non_empty = [r for r in rows if r]
            if not non_empty:
                continue  # 빈 시트
            saw_data_sheet = True
            header = _find_header(non_empty)
            if header is None:
                header_missing = True
                continue
            head_index, time_col, speaker_col, text_col = header
            last_speaker = DEFAULT_SPEAKER
            for row in non_empty[head_index + 1 :]:
                def cell(col: int | None) -> object:
                    return row[col] if col is not None and col < len(row) else None

                text = _text_of(cell(text_col))
                if not text:
                    continue
                if speaker_col is not None:
                    label = _speaker_label(cell(speaker_col))
                    if label:
                        last_speaker = label
                speaker = last_speaker
                stamp = _time_label(cell(time_col)) if time_col is not None else None
                lines.append(f"{stamp} {speaker}: {text}" if stamp else f"{speaker}: {text}")
    finally:
        workbook.close()

    if not saw_data_sheet:
        raise TranscriptTextError("빈 파일입니다(내용이 있는 시트가 없습니다).")
    if not lines:
        if header_missing:
            raise TranscriptTextError(
                "열 이름을 찾지 못했습니다. 첫 행에 '발언'(또는 내용) 열이 필요하고, '시각'과 '화자' 열은 있으면 읽습니다."
            )
        raise TranscriptTextError("읽을 수 있는 발언이 없습니다(발언 열이 비어 있습니다).")
    text = "\n".join(lines)
    if len(text) > settings.transcript_max_chars:
        raise TranscriptTextError(f"글자 수가 너무 많습니다(최대 {settings.transcript_max_chars}자).")
    return text
